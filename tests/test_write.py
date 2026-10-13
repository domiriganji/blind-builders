"""Write and publish: documents, versions, AI suggestions accepted one at a time, export, venues, cover letters."""
import html
import io
import re

import pytest
from fastapi.testclient import TestClient

from studio import db
from studio.a11y_check import check_html
from studio.ai import AIError, ResearchResult, Source
from studio.features import write as write_mod

from .conftest import post, register

BODY = "# Introduction\n\nThe cat sat on teh mat. It was happy.\n\nBlind writers deserve good tools."


def create_doc(client, title="My essay", kind="article", target="") -> int:
    r = post(client, "/write/new", {"title": title, "kind": kind, "target": target}, follow_redirects=False)
    assert r.status_code == 303, r.text
    m = re.fullmatch(r"/write/(\d+)", r.headers["location"])
    assert m, r.headers["location"]
    return int(m.group(1))


def save(client, doc_id, body, title="My essay", target=""):
    return post(client, f"/write/{doc_id}/save", {"body": body, "title": title, "target": target})


def body_of(doc_id) -> str:
    with db.session() as conn:
        return conn.execute("SELECT body FROM write_docs WHERE id = ?", (doc_id,)).fetchone()["body"]


def versions_of(doc_id) -> list[dict]:
    with db.session() as conn:
        return [dict(r) for r in conn.execute("SELECT * FROM write_versions WHERE doc_id = ? ORDER BY id", (doc_id,))]


def suggestion_rows(batch_id) -> list[dict]:
    with db.session() as conn:
        return [dict(r) for r in conn.execute("SELECT * FROM write_suggestions WHERE batch_id = ? ORDER BY number",
                                              (batch_id,))]


def h1(html: str) -> str:
    found = re.findall(r"<h1[^>]*>(.*?)</h1>", html, re.S)
    assert len(found) == 1, found
    return found[0].strip()


def ask_suggestions(client, doc_id, fake_ai, suggestions, note="Mostly clear.", goal="spelling") -> int:
    fake_ai.on("write.suggest", {"overall_note": note, "suggestions": suggestions})
    r = post(client, f"/write/{doc_id}/suggest", {"goal": goal, "instructions": ""}, follow_redirects=False)
    assert r.status_code == 303, r.text
    job_page = client.get(r.headers["location"], follow_redirects=False)
    assert job_page.status_code == 303, job_page.text
    m = re.fullmatch(rf"/write/{doc_id}/suggestions/(\d+)", job_page.headers["location"])
    assert m, job_page.headers["location"]
    return int(m.group(1))


# ----- index and creating ------------------------------------------------------------------------

def test_signed_out_index_explains_and_links_to_sign_in(client):
    r = client.get("/write")
    assert r.status_code == 200
    assert h1(r.text) == "Write and publish"
    assert 'href="/login?next=/write"' in r.text and 'href="/register?next=/write"' in r.text
    assert "never changes your words by itself" in r.text
    assert not check_html(r.text).errors


def test_signed_out_cannot_open_new_form(client):
    r = client.get("/write/new", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"].startswith("/login")


def test_create_and_list(user_client):
    r = user_client.get("/write")
    assert "You have no documents yet" in r.text and "Start a new document" in r.text
    doc_id = create_doc(user_client, "Letter to the council", "letter", "City News")
    page = user_client.get(f"/write/{doc_id}")
    assert page.status_code == 200
    # Straight after creating it, the h1 says what happened; later visits just name the document.
    assert h1(page.text) == "Letter to the council: new document started"
    assert "<title>Letter to the council: new document started, Blind Builders</title>" in page.text
    page = user_client.get(f"/write/{doc_id}")
    assert h1(page.text) == "Letter to the council"
    assert "<title>Letter to the council, Blind Builders</title>" in page.text
    assert "Nothing saved yet" in page.text and "Aiming for: City News" in page.text
    assert 'rows="20"' in page.text and ">Your text" in page.text
    for heading in ["Get suggestions from the AI", "Versions", "Export", "Find where to publish", "Draft a cover letter"]:
        assert f">{heading}</h2>" in page.text
    listing = user_client.get("/write")
    assert f'href="/write/{doc_id}"' in listing.text and "Letter" in listing.text
    assert '<th scope="col">Words</th>' in listing.text and "You have 1 document." in listing.text


def test_new_form_is_prefilled(user_client):
    r = user_client.get("/write/new?title=Grant+story&kind=story")
    assert r.status_code == 200 and h1(r.text) == "Start a new document"
    assert 'value="Grant story"' in r.text
    assert re.search(r'value="story" checked', r.text)
    # An unknown kind falls back to article
    r = user_client.get("/write/new?kind=poem")
    assert re.search(r'value="article" checked', r.text)


def test_create_validation_errors(user_client):
    r = post(user_client, "/write/new", {"title": "", "kind": "poem", "target": "x" * 300})
    assert r.status_code == 400
    assert h1(r.text) == "Please check 3 answers: Start a new document"
    assert 'href="#field-title"' in r.text and 'href="#field-kind"' in r.text and 'href="#field-target"' in r.text
    assert 'aria-invalid="true"' in r.text
    with db.session() as conn:
        assert conn.execute("SELECT COUNT(*) AS n FROM write_docs").fetchone()["n"] == 0


def test_create_requires_csrf(user_client):
    r = user_client.post("/write/new", data={"title": "T", "kind": "article"})
    assert r.status_code == 400


# ----- saving and versions -----------------------------------------------------------------------

def test_save_creates_versions_and_skips_unchanged(user_client):
    doc_id = create_doc(user_client)
    r = save(user_client, doc_id, "First draft.")
    assert r.status_code == 200 and "Saved as version 1" in r.text
    assert "2 words" in r.text
    save(user_client, doc_id, "Second draft, longer.")
    r = save(user_client, doc_id, "Second draft, longer.")
    assert "no new version was needed" in r.text
    versions = versions_of(doc_id)
    assert [v["body"] for v in versions] == ["First draft.", "Second draft, longer."]
    assert all(v["note"] == "Saved by you" for v in versions)
    page = user_client.get(f"/write/{doc_id}/versions")
    assert h1(page.text) == "Versions of My essay"
    assert "2 versions saved" in page.text
    # Newest first
    assert page.text.index("Version 2") < page.text.index("Version 1")
    v1 = user_client.get(f"/write/{doc_id}/versions/{versions[0]['id']}")
    assert h1(v1.text) == "Version 1 of My essay"
    assert "First draft." in v1.text and "Restore version 1" in v1.text
    v2 = user_client.get(f"/write/{doc_id}/versions/{versions[1]['id']}")
    assert "same as your current text" in v2.text and "Restore version" not in v2.text


def test_save_normalises_line_endings_and_updates_title_and_target(user_client):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "Line one\r\nLine two", title="New title", target="The Journal")
    assert body_of(doc_id) == "Line one\nLine two"
    page = user_client.get(f"/write/{doc_id}")
    assert h1(page.text) == "New title" and "Aiming for: The Journal" in page.text


def test_save_validation_errors(user_client):
    doc_id = create_doc(user_client)
    r = save(user_client, doc_id, "Pay with 4111 1111 1111 1111 please", title="")
    assert r.status_code == 400
    assert h1(r.text) == "Please check 2 answers: My essay"
    assert "card number" in r.text and 'href="#field-body"' in r.text and 'href="#field-title"' in r.text
    # The person's text is kept in the form so nothing is lost
    assert "Pay with 4111 1111 1111 1111 please" in r.text
    assert body_of(doc_id) == "" and versions_of(doc_id) == []
    r = save(user_client, doc_id, "x" * (write_mod.MAX_BODY + 1))
    assert r.status_code == 400 and "one document can hold up to 300,000" in r.text


def test_card_number_detector():
    assert write_mod.looks_like_card_number("4111-1111-1111-1111")
    assert write_mod.looks_like_card_number("card 378282246310005 here")
    assert not write_mod.looks_like_card_number("ISBN 978-0-306-40615-7")
    assert not write_mod.looks_like_card_number("1234 5678 9012 3456")  # fails the Luhn check
    assert not write_mod.looks_like_card_number("In 2026 we had 300 members.")


def test_restore_saves_current_first(user_client):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "Old text.")
    save(user_client, doc_id, "New text.")
    first = versions_of(doc_id)[0]
    r = post(user_client, f"/write/{doc_id}/versions/{first['id']}/restore")
    assert r.status_code == 200 and "Restored version 1" in r.text
    assert body_of(doc_id) == "Old text."
    notes = [(v["body"], v["note"]) for v in versions_of(doc_id)]
    assert notes[-1] == ("Old text.", "Restored version 1")
    assert ("New text.", "Saved by you") in notes
    # Restoring a version identical to the current text changes nothing
    latest = versions_of(doc_id)[-1]
    before = len(versions_of(doc_id))
    r = post(user_client, f"/write/{doc_id}/versions/{latest['id']}/restore")
    assert "Nothing was changed" in r.text and len(versions_of(doc_id)) == before


def test_restore_keeps_unversioned_current_text(user_client):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "Version one.")
    first = versions_of(doc_id)[0]
    # Simulate text that was never saved as a version (for example, from before versions existed).
    with db.session() as conn:
        conn.execute("UPDATE write_docs SET body = 'Loose text.' WHERE id = ?", (doc_id,))
    post(user_client, f"/write/{doc_id}/versions/{first['id']}/restore")
    notes = [(v["body"], v["note"]) for v in versions_of(doc_id)]
    assert notes[-2] == ("Loose text.", "Your text before restoring version 1")
    assert notes[-1] == ("Version one.", "Restored version 1")


# ----- suggestions -------------------------------------------------------------------------------

def test_suggestions_keep_only_exact_excerpts_in_document_order(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [
        {"original": "Blind writers deserve good tools.", "proposed": "Blind writers deserve great tools.",
         "reason": "Stronger word."},
        {"original": "teh", "proposed": "the", "reason": "Spelling."},
        {"original": "This sentence is not in the text.", "proposed": "Anything.", "reason": "Invented."},
        {"original": "It was happy.", "proposed": "It was happy.", "reason": "No change."},
    ])
    kind, purpose, kwargs = fake_ai.calls[-1]
    assert kind == "structured" and purpose == "write.suggest"
    assert "teh mat" in kwargs["prompt"] and "Respect the writer's voice" in kwargs["system"]
    # The shared member note is added by AIClient, so the section prompt does not repeat it (docs/VOICE.md 4.18)
    assert "screen reader" not in kwargs["system"] and "never suggest euphemisms" in kwargs["system"]
    rows = suggestion_rows(batch_id)
    assert [(r["number"], r["original"]) for r in rows] == [(1, "teh"), (2, "Blind writers deserve good tools.")]
    assert all(r["status"] == "pending" for r in rows)
    page = user_client.get(f"/write/{doc_id}/suggestions/{batch_id}")
    assert page.status_code == 200
    assert h1(page.text) == "2 suggestions for My essay"
    assert "<title>2 suggestions for My essay, Blind Builders</title>" in page.text
    assert "Suggestion 1 of 2" in page.text and "Suggestion 2 of 2" in page.text
    assert "<strong>Your words:</strong> teh" in page.text and "<strong>Suggested:</strong> the" in page.text
    assert "Was:" not in page.text and "Now:" not in page.text
    text = html.unescape(page.text)
    assert '<strong>Change, computed by the site:</strong> replaces "teh" with "the".' in text
    assert "<strong>The AI's reason:</strong> Spelling." in page.text
    assert ('<strong>Where it is in your text:</strong> In the section "Introduction", paragraph 1: '
            "The cat sat on teh mat. It was happy.") in text
    assert "The AI's overall note" in page.text and "Mostly clear." in page.text and "Written by the AI" in page.text
    assert "2 of the AI's 4 suggestions were left out" in page.text and "Checked by the site" in page.text
    assert "Accept suggestion 1" in page.text and "Reject suggestion 2" in page.text
    assert "Accept all 2 waiting suggestions" in page.text and "Reject all remaining" in page.text
    # Accepting all comes after the last suggestion, so it is never the first button reached
    assert page.text.index("Accept all 2 waiting") > page.text.index("Reject suggestion 2")
    # Nothing changed yet
    assert body_of(doc_id) == BODY
    # The document page links the batch
    assert f"/write/{doc_id}/suggestions/{batch_id}" in user_client.get(f"/write/{doc_id}").text


def test_one_suggestion_heading_is_singular(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [{"original": "teh", "proposed": "the", "reason": "Spelling."}])
    page = user_client.get(f"/write/{doc_id}/suggestions/{batch_id}")
    assert h1(page.text) == "1 suggestion for My essay"
    assert "each suggestion quotes your text word for word" in page.text


def test_no_suggestions(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [], note="This reads well.")
    page = user_client.get(f"/write/{doc_id}/suggestions/{batch_id}")
    assert h1(page.text) == "No suggestions for My essay"
    assert "This reads well." in page.text and "Accept all" not in page.text


def test_accept_applies_exactly_once(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "teh cat and teh dog")
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [{"original": "teh cat", "proposed": "the cat",
                                                               "reason": "Spelling."}])
    s = suggestion_rows(batch_id)[0]
    r = post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{s['id']}/accept")
    assert r.status_code == 200 and "Suggestion 1 accepted" in r.text
    assert body_of(doc_id) == "the cat and teh dog"
    assert versions_of(doc_id)[-1]["note"] == "Accepted suggestion 1"
    assert versions_of(doc_id)[-1]["body"] == "the cat and teh dog"
    # Accepting again does nothing more
    r = post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{s['id']}/accept")
    assert "already decided" in r.text
    assert body_of(doc_id) == "the cat and teh dog"
    assert len(versions_of(doc_id)) == 2
    assert suggestion_rows(batch_id)[0]["status"] == "accepted"
    assert "Accepted. Your text has this change." in r.text


def test_stale_suggestion_is_detected(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [{"original": "teh", "proposed": "the", "reason": "Spelling."}])
    save(user_client, doc_id, BODY.replace("teh", "thee"))
    page = user_client.get(f"/write/{doc_id}/suggestions/{batch_id}")
    assert "This no longer matches your text" in page.text
    s = suggestion_rows(batch_id)[0]
    r = post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{s['id']}/accept")
    assert "no longer matches your text" in r.text and "Nothing was changed" in r.text
    assert suggestion_rows(batch_id)[0]["status"] == "stale"
    assert body_of(doc_id) == BODY.replace("teh", "thee")
    assert "Out of date." in r.text


def test_reject_changes_nothing(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [{"original": "teh", "proposed": "the", "reason": "Spelling."}])
    s = suggestion_rows(batch_id)[0]
    r = post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{s['id']}/reject")
    assert "Suggestion 1 rejected" in r.text
    assert suggestion_rows(batch_id)[0]["status"] == "rejected"
    assert body_of(doc_id) == BODY and len(versions_of(doc_id)) == 1
    # A rejected suggestion cannot then be accepted
    post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{s['id']}/accept")
    assert body_of(doc_id) == BODY


def test_accept_all_that_still_match(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "Alpha one. Beta two. Gamma three.")
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [
        {"original": "Alpha one.", "proposed": "Alpha 1.", "reason": "Digits."},
        {"original": "Beta two.", "proposed": "Beta 2.", "reason": "Digits."},
        {"original": "Gamma three.", "proposed": "Gamma 3.", "reason": "Digits."},
    ])
    rows = suggestion_rows(batch_id)
    post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{rows[1]['id']}/reject")
    # The person edits the text so suggestion 3 no longer matches
    save(user_client, doc_id, "Alpha one. Beta two. Gamma 3!")
    r = post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/accept-all")
    assert r.status_code == 200
    assert "Accepted 1 suggestion" in r.text and "1 suggestion no longer matched your text" in r.text
    assert body_of(doc_id) == "Alpha 1. Beta two. Gamma 3!"
    assert [s["status"] for s in suggestion_rows(batch_id)] == ["accepted", "rejected", "stale"]
    assert versions_of(doc_id)[-1]["note"] == "Accepted suggestion 1"
    # Nothing left to do
    r = post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/accept-all")
    assert "no suggestions waiting" in r.text


def test_accept_all_several_saves_one_version(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "teh one, teh two.")
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [
        {"original": "teh one", "proposed": "the one", "reason": "Spelling."},
        {"original": "teh two", "proposed": "the two", "reason": "Spelling."},
    ])
    post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/accept-all")
    assert body_of(doc_id) == "the one, the two."
    assert versions_of(doc_id)[-1]["note"] == "Accepted suggestions 1 and 2"
    assert len(versions_of(doc_id)) == 2


def test_reject_all_remaining(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [
        {"original": "teh", "proposed": "the", "reason": "Spelling."},
        {"original": "It was happy.", "proposed": "It was glad.", "reason": "Variety."},
    ])
    first = suggestion_rows(batch_id)[0]
    post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{first['id']}/accept")
    r = post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/reject-all")
    assert "Rejected 1 suggestion" in r.text
    assert [s["status"] for s in suggestion_rows(batch_id)] == ["accepted", "rejected"]
    assert body_of(doc_id) == BODY.replace("teh", "the", 1)
    assert "waiting suggestion" not in r.text


def test_suggest_validation(user_client, fake_ai):
    doc_id = create_doc(user_client)
    r = post(user_client, f"/write/{doc_id}/suggest", {"goal": "spelling"})
    assert r.status_code == 400 and "has no text yet" in r.text
    assert h1(r.text) == "Please check 1 answer: My essay"
    save(user_client, doc_id, BODY)
    r = post(user_client, f"/write/{doc_id}/suggest", {"goal": "nonsense"})
    assert r.status_code == 400 and 'href="#field-goal"' in r.text
    r = post(user_client, f"/write/{doc_id}/suggest", {"goal": "publication"})
    assert r.status_code == 400 and "Fill in the publication you are aiming for" in r.text
    assert 'href="#field-target"' in r.text
    r = post(user_client, f"/write/{doc_id}/suggest", {"goal": "clearer", "instructions": "y" * 3000})
    assert r.status_code == 400 and 'href="#field-instructions"' in r.text
    assert fake_ai.calls == []


def test_suggest_for_publication_sends_target_and_instructions(user_client, fake_ai):
    doc_id = create_doc(user_client, target="Disability Studies Quarterly")
    save(user_client, doc_id, BODY, target="Disability Studies Quarterly")
    fake_ai.on("write.suggest", {"overall_note": "", "suggestions": []})
    post(user_client, f"/write/{doc_id}/suggest", {"goal": "publication", "instructions": "Keep British spelling"})
    prompt = fake_ai.calls[-1][2]["prompt"]
    assert "Disability Studies Quarterly" in prompt and "Keep British spelling" in prompt


def test_suggest_ai_error_is_shown_on_job_page(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    fake_ai.on("write.suggest", AIError("The AI is busy right now (rate limit). Please try again in a minute."))
    r = post(user_client, f"/write/{doc_id}/suggest", {"goal": "spelling"})
    assert "This did not finish" in r.text and "rate limit" in r.text


def test_where_in_text_leads_with_the_place():
    body = "word " * 30 + "teh target " + "more " * 30
    ctx = write_mod.where_in_text(body, "teh")
    assert ctx.startswith("Paragraph 1: ...word ") and ctx.endswith(" more...") and "teh target" in ctx
    assert write_mod.where_in_text("teh", "teh") == "Paragraph 1: teh"
    assert write_mod.where_in_text("abc", "zzz") == ""
    # The place comes first, Markdown marks are left out, and the words stay inside their paragraph
    body = "# Abstract\n\nFirst paragraph.\n\nWe assume teh reader.\n\n## Methods\n\nWe paid teh fee."
    assert write_mod.where_in_text(body, "assume teh") == 'In the section "Abstract", paragraph 2: We assume teh reader.'
    assert write_mod.where_in_text(body, "paid teh") == 'In the section "Methods", paragraph 1: We paid teh fee.'
    assert write_mod.where_in_text(body, "Methods") == 'In the heading "Methods".'


def test_describe_change_is_worked_out_word_by_word():
    describe = write_mod.describe_change
    assert describe("teh", "the") == 'replaces "teh" with "the".'
    assert describe("We assume the reader knows.", "We assume that the reader knows.") == 'adds "that" after "assume".'
    assert describe("It was very happy.", "It was happy.") == 'removes "very".'
    assert describe("It was happy.", "") == 'removes "It was happy."'
    assert describe("Go now", "") == 'removes "Go now".'
    assert describe("Start here.", "Now Start here.") == 'adds "Now" before "Start".'
    assert describe("a  b", "a b") == "changes only spaces or line breaks."
    assert describe("one two three four five six seven eight", "1 two 3 four 5 six 7 eight") == (
        "makes 4 changes in this passage.")


def test_filter_suggestions_unit():
    body = "one two three two"
    kept, dropped = write_mod.filter_suggestions(body, [
        {"original": "three", "proposed": "3", "reason": "a"},
        {"original": "one", "proposed": "1", "reason": "b"},
        {"original": "", "proposed": "x", "reason": "empty"},
        {"original": "four", "proposed": "4", "reason": "missing"},
        {"original": "one", "proposed": "1", "reason": "duplicate"},
    ])
    assert [k["original"] for k in kept] == ["one", "three"]
    assert dropped == 3


# ----- export ------------------------------------------------------------------------------------

EXPORT_BODY = "# Methods\n\nWe paid $5 & saved 10% on item #3_b {x}.\n\n## Results\n\n- first point\n- second point\n\nA <script>alert(1)</script> tag."


def test_export_page_lists_formats(user_client):
    doc_id = create_doc(user_client, "Cost & value")
    page = user_client.get(f"/write/{doc_id}/export")
    assert h1(page.text) == "Export Cost &amp; value"
    for ext in ["docx", "md", "txt", "html", "tex"]:
        assert f'href="/write/{doc_id}/export.{ext}"' in page.text
    assert not check_html(page.text).errors


def test_export_markdown_and_text(user_client):
    doc_id = create_doc(user_client, "Cost study")
    save(user_client, doc_id, EXPORT_BODY, title="Cost study")
    md = user_client.get(f"/write/{doc_id}/export.md")
    assert md.status_code == 200 and md.headers["content-type"].startswith("text/markdown")
    assert 'attachment; filename="cost-study.md"' == md.headers["content-disposition"]
    # The title is the only level 1 heading; the member's headings follow from level 2.
    assert md.text.startswith("# Cost study\n\n## Methods")
    assert "\n### Results\n" in md.text and md.text.count("\n# ") == 0
    txt = user_client.get(f"/write/{doc_id}/export.txt")
    assert txt.headers["content-type"].startswith("text/plain")
    assert txt.text.startswith("Cost study\n\n") and "We paid $5" in txt.text


def test_export_skips_repeated_title_heading(user_client):
    doc_id = create_doc(user_client, "Cost study")
    save(user_client, doc_id, "# Cost study\n\nBody text.", title="Cost study")
    assert user_client.get(f"/write/{doc_id}/export.md").text == "# Cost study\n\nBody text.\n"


def test_export_html_is_complete_and_accessible(user_client):
    doc_id = create_doc(user_client, "Cost study")
    save(user_client, doc_id, EXPORT_BODY, title="Cost study")
    r = user_client.get(f"/write/{doc_id}/export.html")
    assert r.headers["content-type"].startswith("text/html")
    assert r.headers["content-disposition"].startswith("attachment")
    html = r.text
    assert html.startswith("<!doctype html>") and '<html lang="en">' in html
    assert "<title>Cost study</title>" in html
    assert h1(html) == "Cost study"
    assert "<h2>Methods</h2>" in html and "<h3>Results</h3>" in html
    assert "<script>" not in html and "&lt;script&gt;" in html
    report = check_html(html)
    assert not report.errors, report.findings


def test_export_docx_parses_with_headings(user_client):
    from docx import Document
    doc_id = create_doc(user_client, "Cost study")
    save(user_client, doc_id, EXPORT_BODY, title="Cost study")
    r = user_client.get(f"/write/{doc_id}/export.docx")
    assert r.status_code == 200
    assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    document = Document(io.BytesIO(r.content))
    paragraphs = [(p.style.name, p.text) for p in document.paragraphs]
    assert paragraphs[0] == ("Heading 1", "Cost study")
    assert ("Heading 2", "Methods") in paragraphs and ("Heading 3", "Results") in paragraphs
    assert ("List Bullet", "first point") in paragraphs
    assert any(t == "We paid $5 & saved 10% on item #3_b {x}." for _, t in paragraphs)
    assert not any(t.startswith("#") for _, t in paragraphs)
    assert document.core_properties.title == "Cost study"


def test_export_tex_escapes_special_characters(user_client):
    doc_id = create_doc(user_client, "Cost & 100% study_1")
    save(user_client, doc_id, EXPORT_BODY, title="Cost & 100% study_1")
    r = user_client.get(f"/write/{doc_id}/export.tex")
    assert r.status_code == 200 and r.headers["content-disposition"].endswith('.tex"')
    tex = r.text
    assert tex.startswith("\\documentclass{article}") and tex.rstrip().endswith("\\end{document}")
    assert "\\title{Cost \\& 100\\% study\\_1}" in tex
    assert "\\section{Methods}" in tex and "\\subsection{Results}" in tex
    assert "We paid \\$5 \\& saved 10\\% on item \\#3\\_b \\{x\\}." in tex
    assert write_mod.latex_escape("a\\b~c^d") == "a\\textbackslash{}b\\textasciitilde{}c\\textasciicircum{}d"
    body = tex.split("\\maketitle", 1)[1]
    assert "$5" not in body.replace("\\$5", "") and "#3" not in body.replace("\\#3", "")


def test_export_empty_document(user_client):
    doc_id = create_doc(user_client, "Empty")
    from docx import Document
    assert Document(io.BytesIO(user_client.get(f"/write/{doc_id}/export.docx").content)).paragraphs[0].text == "Empty"
    assert "\\begin{document}" in user_client.get(f"/write/{doc_id}/export.tex").text
    assert check_html(user_client.get(f"/write/{doc_id}/export.html").text).errors == []


# ----- where to publish and cover letters --------------------------------------------------------

def test_venues_research_with_sources(user_client, fake_ai):
    doc_id = create_doc(user_client, "Screen readers in class", "paper")
    save(user_client, doc_id, BODY, title="Screen readers in class")
    fake_ai.on("write.venues", ResearchResult(
        text="Summary: two good fits.\n\n# Journal of Examples\n\nNo fee found. Nothing found about screen reader access.",
        sources=[Source(url="https://example.org/guidelines", title="Author guidelines", quote="Submit online."),
                 Source(url="javascript:alert(1)", title="Bad link")]))
    r = post(user_client, f"/write/{doc_id}/venues")
    assert r.status_code == 200
    assert h1(r.text) == "Where to publish: Screen readers in class"
    assert "<h3>Journal of Examples</h3>" in r.text and ">Written by the AI (not checked)</h2>" in r.text
    assert 'href="https://example.org/guidelines" rel="noopener noreferrer"' in r.text
    assert "javascript:" not in r.text
    assert "Sources (1)" in r.text
    kind, purpose, kwargs = fake_ai.calls[-1]
    assert kind == "research" and purpose == "write.venues"
    assert "open access fee" in kwargs["prompt"] and "screen reader" in kwargs["prompt"]
    assert not check_html(r.text).errors
    with db.session() as conn:
        rid = conn.execute("SELECT id FROM write_research WHERE doc_id = ?", (doc_id,)).fetchone()["id"]
    assert f'href="/write/{doc_id}/publish/{rid}"' in user_client.get(f"/write/{doc_id}").text


def test_venues_without_sources_warns(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    fake_ai.on("write.venues", ResearchResult(text="Nothing reliable found.", sources=[]))
    r = post(user_client, f"/write/{doc_id}/venues")
    assert "did not list any sources" in r.text


def test_venues_needs_text(user_client, fake_ai):
    doc_id = create_doc(user_client)
    r = post(user_client, f"/write/{doc_id}/venues")
    assert "has no text yet" in r.text and fake_ai.calls == []


def test_cover_letter_creates_new_document(user_client, fake_ai):
    doc_id = create_doc(user_client, "Screen readers in class", "paper", "Journal of Examples")
    save(user_client, doc_id, BODY, title="Screen readers in class", target="Journal of Examples")
    page = user_client.get(f"/write/{doc_id}")
    assert ">Save and draft a cover letter for Journal of Examples</button>" in page.text
    fake_ai.on("write.cover_letter", "Dear [Editor's name],\n\nI am submitting my paper.\n\n[Your name]")
    r = post(user_client, f"/write/{doc_id}/cover-letter", {"notes": "I teach braille."}, follow_redirects=False)
    job = user_client.get(r.headers["location"], follow_redirects=False)
    new_url = job.headers["location"]
    new_id = int(new_url.rsplit("/", 1)[1])
    assert new_id != doc_id
    with db.session() as conn:
        new = dict(conn.execute("SELECT * FROM write_docs WHERE id = ?", (new_id,)).fetchone())
    assert new["kind"] == "letter" and new["title"] == "Cover letter: Screen readers in class"
    assert new["target"] == "Journal of Examples" and new["body"].startswith("Dear [Editor's name]")
    assert [(v["note"], v["body"]) for v in versions_of(new_id)] == [("Drafted by the AI", new["body"])]
    page = user_client.get(new_url)
    assert h1(page.text) == "Cover letter: Screen readers in class"
    assert "Drafted by the AI" in page.text and "The AI wrote this draft" in page.text
    assert "change anything that is not true before you send it." in page.text
    kind, purpose, kwargs = fake_ai.calls[-1]
    assert kind == "text" and purpose == "write.cover_letter"
    assert "Journal of Examples" in kwargs["prompt"] and "I teach braille." in kwargs["prompt"]
    # The original is unchanged
    assert body_of(doc_id) == BODY


AI_DRAFT_PANEL = "The AI wrote this draft. Nobody has checked it. Fill in everything in square brackets"


def test_the_ai_draft_panel_shows_for_any_ai_drafted_last_version(user_client):
    doc_id = create_doc(user_client, "Exams and screen readers", "paper")
    save(user_client, doc_id, BODY, title="Exams and screen readers")
    assert AI_DRAFT_PANEL not in user_client.get(f"/write/{doc_id}").text
    # Research copies its drafts into Write with a longer note that starts the same way.
    with db.session() as conn:
        conn.execute("INSERT INTO write_versions (doc_id, body, note, created_at) VALUES (?, ?, ?, ?)",
                     (doc_id, BODY, "Drafted by the AI in Research (not checked)", db.now()))
    page = user_client.get(f"/write/{doc_id}")
    assert page.status_code == 200
    assert '<p class="panel">' + AI_DRAFT_PANEL + ", and change anything that is not true before you share it.</p>" \
        in page.text
    # Once the member saves their own version, the last version is theirs and the panel goes.
    save(user_client, doc_id, BODY + "\n\nMy own words.", title="Exams and screen readers")
    assert AI_DRAFT_PANEL not in user_client.get(f"/write/{doc_id}").text


def test_cover_letter_needs_target(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    page = user_client.get(f"/write/{doc_id}")
    assert 'href="#field-target">Go to the publication field' in page.text
    assert 'formaction="/write/%d/cover-letter">Save and draft a cover letter</button>' % doc_id in page.text
    r = post(user_client, f"/write/{doc_id}/cover-letter", {"notes": ""})
    assert r.status_code == 400 and "Fill in the publication you are sending this to" in r.text
    # The error summary links to the field that fixes the problem, which is a real input.
    assert 'href="#field-target"' in r.text and '<input id="field-target"' in r.text
    assert fake_ai.calls == []


# ----- deleting ----------------------------------------------------------------------------------

def test_delete_document(user_client, fake_ai):
    doc_id = create_doc(user_client, "Throwaway")
    save(user_client, doc_id, BODY, title="Throwaway")
    ask_suggestions(user_client, doc_id, fake_ai, [{"original": "teh", "proposed": "the", "reason": "Spelling."}])
    page = user_client.get(f"/write/{doc_id}/delete")
    assert h1(page.text) == "Delete Throwaway?" and "cannot be undone" in page.text
    r = post(user_client, f"/write/{doc_id}/delete")
    assert r.status_code == 200 and "Deleted" in r.text
    assert user_client.get(f"/write/{doc_id}").status_code == 404
    with db.session() as conn:
        for table in ["write_versions", "write_suggestions", "write_batches", "write_research"]:
            assert conn.execute(f"SELECT COUNT(*) AS n FROM {table}").fetchone()["n"] == 0


# ----- ownership ---------------------------------------------------------------------------------

def test_other_members_get_404_everywhere(client, fake_ai):
    register(client, "owner@example.com", "Owner")
    doc_id = create_doc(client, target="Some Journal")
    save(client, doc_id, BODY, target="Some Journal")
    batch_id = ask_suggestions(client, doc_id, fake_ai, [{"original": "teh", "proposed": "the", "reason": "Spelling."}])
    sid = suggestion_rows(batch_id)[0]["id"]
    vid = versions_of(doc_id)[0]["id"]
    fake_ai.on("write.venues", ResearchResult(text="x", sources=[]))
    post(client, f"/write/{doc_id}/venues")
    with db.session() as conn:
        rid = conn.execute("SELECT id FROM write_research").fetchone()["id"]
    post(client, "/logout")
    register(client, "other@example.com", "Other")
    calls_before = len(fake_ai.calls)
    gets = ["", "/versions", f"/versions/{vid}", "/export", "/export.md", "/export.txt", "/export.html",
            "/export.docx", "/export.tex", f"/suggestions/{batch_id}", f"/publish/{rid}", "/delete"]
    for suffix in gets:
        r = client.get(f"/write/{doc_id}{suffix}")
        assert r.status_code == 404, suffix
    posts = [("/save", {"body": "hacked", "title": "x"}), ("/suggest", {"goal": "spelling"}), ("/venues", {}),
             ("/cover-letter", {}), (f"/versions/{vid}/restore", {}), (f"/suggestions/{batch_id}/{sid}/accept", {}),
             (f"/suggestions/{batch_id}/{sid}/reject", {}), (f"/suggestions/{batch_id}/accept-all", {}),
             (f"/suggestions/{batch_id}/reject-all", {}), ("/delete", {})]
    for suffix, data in posts:
        r = post(client, f"/write/{doc_id}{suffix}", data)
        assert r.status_code == 404, suffix
    assert len(fake_ai.calls) == calls_before
    assert body_of(doc_id) == BODY
    assert suggestion_rows(batch_id)[0]["status"] == "pending"
    assert "My essay" not in client.get("/write").text


def test_suggestion_and_version_ids_must_belong_to_the_document(user_client, fake_ai):
    a = create_doc(user_client, "A")
    b = create_doc(user_client, "B")
    save(user_client, a, BODY, title="A")
    save(user_client, b, BODY, title="B")
    batch_a = ask_suggestions(user_client, a, fake_ai, [{"original": "teh", "proposed": "the", "reason": "Spelling."}])
    sid = suggestion_rows(batch_a)[0]["id"]
    vid_a = versions_of(a)[0]["id"]
    assert user_client.get(f"/write/{b}/suggestions/{batch_a}").status_code == 404
    assert post(user_client, f"/write/{b}/suggestions/{batch_a}/{sid}/accept").status_code == 404
    assert user_client.get(f"/write/{b}/versions/{vid_a}").status_code == 404
    assert post(user_client, f"/write/{b}/versions/{vid_a}/restore").status_code == 404
    assert user_client.get("/write/999999").status_code == 404
    assert body_of(a) == BODY and body_of(b) == BODY


# ----- no AI available ---------------------------------------------------------------------------

def test_no_ai_explains_what_to_do(no_ai_app):
    client = TestClient(no_ai_app)
    register(client)
    doc_id = create_doc(client, target="Some Journal")
    save(client, doc_id, BODY, target="Some Journal")
    for url, data in [(f"/write/{doc_id}/suggest", {"goal": "spelling"}), (f"/write/{doc_id}/venues", {}),
                      (f"/write/{doc_id}/cover-letter", {"notes": ""})]:
        r = post(client, url, data, follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == f"/write/{doc_id}"
        page = client.get(f"/write/{doc_id}")
        assert "connect your own Claude API key" in page.text
    # Writing, versions and export still work without the AI
    assert client.get(f"/write/{doc_id}/export.docx").status_code == 200
    assert len(versions_of(doc_id)) == 1
    with db.session() as conn:
        assert conn.execute("SELECT COUNT(*) AS n FROM jobs").fetchone()["n"] == 0


# ----- every page passes the accessibility checker -----------------------------------------------

def test_all_write_pages_pass_checks(user_client, fake_ai):
    doc_id = create_doc(user_client, "Checked paper", "paper", "Some Journal")
    save(user_client, doc_id, BODY, title="Checked paper", target="Some Journal")
    save(user_client, doc_id, BODY + "\n\nMore.", title="Checked paper", target="Some Journal")
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [
        {"original": "teh", "proposed": "the", "reason": "Spelling."},
        {"original": "It was happy.", "proposed": "", "reason": "Not needed."},
        {"original": "missing words", "proposed": "x", "reason": "Dropped."},
    ], note="# Overall\n\nGood work.")
    first = suggestion_rows(batch_id)[0]
    post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{first['id']}/accept")
    fake_ai.on("write.venues", ResearchResult(text="# Place\n\nDetails.", sources=[Source(url="https://example.org", title="Ex")]))
    post(user_client, f"/write/{doc_id}/venues")
    with db.session() as conn:
        rid = conn.execute("SELECT id FROM write_research").fetchone()["id"]
    vid = versions_of(doc_id)[0]["id"]
    paths = ["/write", "/write/new", f"/write/{doc_id}", f"/write/{doc_id}/versions", f"/write/{doc_id}/versions/{vid}",
             f"/write/{doc_id}/export", f"/write/{doc_id}/suggestions/{batch_id}", f"/write/{doc_id}/publish/{rid}",
             f"/write/{doc_id}/delete"]
    problems = []
    for path in paths:
        r = user_client.get(path)
        assert r.status_code == 200, path
        h1(r.text)
        for f in check_html(r.text, name=path).findings:
            problems.append(f"{path}: {f.severity} {f.rule}: {f.message} {f.where}")
    assert not problems, "\n".join(problems)
    # Every page is reachable from the section index
    doc_page = user_client.get(f"/write/{doc_id}").text
    for path in paths[3:]:
        if path in (f"/write/{doc_id}/versions/{vid}",):
            assert path in user_client.get(f"/write/{doc_id}/versions").text
        else:
            assert f'href="{path}"' in doc_page, path
    assert f'href="/write/{doc_id}"' in user_client.get("/write").text


def test_validation_error_pages_pass_checks(user_client):
    r = post(user_client, "/write/new", {"title": "", "kind": ""})
    assert not check_html(r.text).errors
    doc_id = create_doc(user_client)
    r = post(user_client, f"/write/{doc_id}/suggest", {"goal": "spelling"})
    assert not check_html(r.text).errors
    r = post(user_client, f"/write/{doc_id}/cover-letter", {})
    assert not check_html(r.text).errors


@pytest.mark.parametrize("title,expected", [("Cost & value!", "cost-value"), ("***", "document"), ("Été 2026", "ete-2026")])
def test_slug(title, expected):
    assert write_mod._slug(title) == expected


# ----- regression tests from the audit -----------------------------------------------------------

def test_h1_says_what_just_happened_once(user_client, fake_ai):
    doc_id = create_doc(user_client)
    r = save(user_client, doc_id, BODY)
    assert h1(r.text) == "My essay: saved as version 1" and "<title>My essay: saved as version 1, Blind Builders</title>" in r.text
    assert h1(user_client.get(f"/write/{doc_id}").text) == "My essay"  # only once
    r = save(user_client, doc_id, BODY)
    assert h1(r.text) == "My essay: saved, the text had not changed"
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [{"original": "teh", "proposed": "the", "reason": "Spelling."},
                                                              {"original": "It was happy.", "proposed": "Happy.", "reason": "Short."}])
    rows = suggestion_rows(batch_id)
    r = post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{rows[0]['id']}/accept")
    assert h1(r.text) == "2 suggestions for My essay: suggestion 1 accepted"
    r = post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{rows[1]['id']}/reject")
    assert h1(r.text) == "2 suggestions for My essay: suggestion 2 rejected"
    r = post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/accept-all")
    assert h1(r.text) == "2 suggestions for My essay: nothing was waiting, nothing changed"
    first = versions_of(doc_id)[0]
    r = post(user_client, f"/write/{doc_id}/versions/{first['id']}/restore")
    assert h1(r.text) == "My essay: version 1 restored"
    r = post(user_client, f"/write/{doc_id}/delete")
    assert h1(r.text) == "Write and publish: My essay deleted"
    assert h1(user_client.get("/write").text) == "Write and publish"


def test_outcome_is_not_shown_on_another_page_or_when_old(user_client, monkeypatch):
    a = create_doc(user_client, "A")
    b = create_doc(user_client, "B")
    post(user_client, f"/write/{a}/save", {"body": "Text.", "title": "A"}, follow_redirects=False)
    # The outcome was for A; opening B first throws it away
    assert h1(user_client.get(f"/write/{b}").text) == "B"
    assert h1(user_client.get(f"/write/{a}").text) == "A"
    post(user_client, f"/write/{a}/save", {"body": "Text 2.", "title": "A"}, follow_redirects=False)
    real_time = write_mod.time.time
    monkeypatch.setattr(write_mod.time, "time", lambda: real_time() + write_mod.OUTCOME_SECONDS + 5)
    assert h1(user_client.get(f"/write/{a}").text) == "A"


def test_double_accept_applies_the_change_only_once(user_client, fake_ai, monkeypatch):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "teh cat and teh dog")
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [{"original": "teh cat", "proposed": "the cat",
                                                               "reason": "Spelling."}])
    s = suggestion_rows(batch_id)[0]
    post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{s['id']}/accept")
    # A second press that was already on its way, having read the suggestion while it was still pending
    monkeypatch.setattr(write_mod, "_owned_suggestion", lambda doc, b, sid: dict(s))
    r = post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{s['id']}/accept")
    assert "already decided" in r.text
    assert body_of(doc_id) == "the cat and teh dog"
    assert len(versions_of(doc_id)) == 2
    r = post(user_client, f"/write/{doc_id}/suggestions/{batch_id}/{s['id']}/reject")
    assert "already decided" in r.text and suggestion_rows(batch_id)[0]["status"] == "accepted"


def test_double_save_and_double_restore_keep_one_version(user_client, monkeypatch):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "One.")
    save(user_client, doc_id, "Two.")
    real_owned = write_mod._owned_doc
    # Second presses that read the document before the first press had saved
    monkeypatch.setattr(write_mod, "_owned_doc", lambda d, u: dict(real_owned(d, u), body="One."))
    r = save(user_client, doc_id, "Two.")
    assert "no new version was needed" in r.text and len(versions_of(doc_id)) == 2
    monkeypatch.setattr(write_mod, "_owned_doc", real_owned)
    first = versions_of(doc_id)[0]
    post(user_client, f"/write/{doc_id}/versions/{first['id']}/restore")
    count = len(versions_of(doc_id))
    monkeypatch.setattr(write_mod, "_owned_doc", lambda d, u: dict(real_owned(d, u), body="Two."))
    r = post(user_client, f"/write/{doc_id}/versions/{first['id']}/restore")
    assert "same as your current text" in r.text and len(versions_of(doc_id)) == count
    assert body_of(doc_id) == "One."


def test_pressing_an_ai_button_twice_does_not_start_or_pay_twice(user_client, fake_ai):
    doc_id = create_doc(user_client, target="Some Journal")
    save(user_client, doc_id, BODY, target="Some Journal")
    with db.session() as conn:
        conn.execute("INSERT INTO jobs (id, user_id, kind, title, status, created_at) VALUES (?, ?, ?, ?, 'running', ?)",
                     ("runningjob1", 1, f"write.suggest:{doc_id}", "Suggestions for My essay", db.now()))
    r = post(user_client, f"/write/{doc_id}/suggest", {"goal": "spelling"}, follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/jobs/runningjob1"
    assert fake_ai.calls == []
    assert "already working on this" in user_client.get("/jobs/runningjob1").text
    # Venues and cover letters are separate kinds of work, so they still start
    fake_ai.on("write.venues", ResearchResult(text="Found one.", sources=[]))
    assert post(user_client, f"/write/{doc_id}/venues").status_code == 200
    assert fake_ai.purposes() == ["write.venues"]
    # A job left "running" for a long time (for example after a restart) does not block new work
    with db.session() as conn:
        conn.execute("UPDATE jobs SET created_at = '2020-01-01T00:00:00+00:00' WHERE id = 'runningjob1'")
    fake_ai.on("write.suggest", {"overall_note": "", "suggestions": []})
    r = post(user_client, f"/write/{doc_id}/suggest", {"goal": "spelling"}, follow_redirects=False)
    assert r.headers["location"] != "/jobs/runningjob1" and "write.suggest" in fake_ai.purposes()


def test_filter_suggestions_drops_missing_or_wrongly_typed_fields():
    body = "None of 5 cats sat."
    kept, dropped = write_mod.filter_suggestions(body, [
        {"original": "cats", "reason": "proposed is missing"},          # must not become "remove these words"
        {"original": "cats", "proposed": None, "reason": "null"},
        {"original": None, "proposed": "x", "reason": "would be the text 'None'"},
        {"original": 5, "proposed": "five", "reason": "a number, not text"},
        {"original": "cats", "proposed": "x" * (write_mod.MAX_PROPOSED + 1), "reason": "too long"},
        "not an object",
        {"original": "sat", "proposed": "sat down", "reason": ["not", "text"]},
        {"original": "cats", "proposed": "", "reason": "  Remove.  "},
    ])
    assert kept == [{"original": "cats", "proposed": "", "reason": "Remove."},
                    {"original": "sat", "proposed": "sat down", "reason": ""}]
    assert dropped == 6


@pytest.mark.parametrize("answer", [{"overall_note": "Looks fine."}, {"overall_note": "x", "suggestions": "none"},
                                    ["not", "an", "object"]])
def test_incomplete_ai_answer_is_not_shown_as_no_suggestions(user_client, fake_ai, answer):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    fake_ai.on("write.suggest", answer)
    r = post(user_client, f"/write/{doc_id}/suggest", {"goal": "spelling"})
    assert "This did not finish" in r.text and "incomplete" in r.text
    with db.session() as conn:
        assert conn.execute("SELECT COUNT(*) AS n FROM write_batches").fetchone()["n"] == 0


def test_overall_note_of_wrong_type_is_left_out(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [{"original": "teh", "proposed": "the"}], note=None)
    page = user_client.get(f"/write/{doc_id}/suggestions/{batch_id}")
    assert "The AI did not write an overall note." in page.text and "The AI gave no reason." in page.text


def test_document_deleted_while_the_ai_works(user_client, fake_ai):
    doc_id = create_doc(user_client, target="Some Journal")
    save(user_client, doc_id, BODY, target="Some Journal")

    def delete_then(value):
        def answer(**kwargs):
            with db.session() as conn:
                conn.execute("DELETE FROM write_docs WHERE id = ?", (doc_id,))
            return value
        return answer

    for purpose, url, value in [
        ("write.suggest", "/suggest", {"overall_note": "", "suggestions": [{"original": "teh", "proposed": "the", "reason": ""}]}),
        ("write.venues", "/venues", ResearchResult(text="Found.", sources=[])),
        ("write.cover_letter", "/cover-letter", "Dear [Editor's name],"),
    ]:
        with db.session() as conn:
            conn.execute("INSERT OR REPLACE INTO write_docs (id, user_id, title, kind, body, target, created_at, updated_at)"
                         " VALUES (?, 1, 'My essay', 'article', ?, 'Some Journal', ?, ?)", (doc_id, BODY, db.now(), db.now()))
        fake_ai.on(purpose, delete_then(value))
        r = post(user_client, f"/write/{doc_id}{url}", {"goal": "spelling"})
        assert "This did not finish" in r.text, purpose
        assert "deleted while the AI was working" in r.text, purpose
    with db.session() as conn:
        for table in ["write_batches", "write_research"]:
            assert conn.execute(f"SELECT COUNT(*) AS n FROM {table}").fetchone()["n"] == 0
        assert conn.execute("SELECT COUNT(*) AS n FROM write_docs").fetchone()["n"] == 0


def test_empty_ai_answers_save_nothing(user_client, fake_ai):
    doc_id = create_doc(user_client, target="Some Journal")
    save(user_client, doc_id, BODY, target="Some Journal")
    fake_ai.on("write.venues", ResearchResult(text="   ", sources=[Source(url="https://example.org", title="Ex")]))
    r = post(user_client, f"/write/{doc_id}/venues")
    assert "This did not finish" in r.text and "did not write an answer" in r.text
    fake_ai.on("write.cover_letter", "  \n ")
    r = post(user_client, f"/write/{doc_id}/cover-letter", {"notes": ""})
    assert "This did not finish" in r.text and "did not write a letter" in r.text
    with db.session() as conn:
        assert conn.execute("SELECT COUNT(*) AS n FROM write_research").fetchone()["n"] == 0
        assert conn.execute("SELECT COUNT(*) AS n FROM write_docs").fetchone()["n"] == 1


def test_errors_point_at_the_field_that_fixes_them(user_client, fake_ai):
    doc_id = create_doc(user_client)
    r = post(user_client, f"/write/{doc_id}/venues")
    assert r.status_code == 400 and h1(r.text) == "Please check 1 answer: My essay"
    assert 'href="#field-body"' in r.text and "has no text yet" in r.text
    r = post(user_client, f"/write/{doc_id}/cover-letter", {"notes": "n" * 2001})
    assert r.status_code == 400
    for field in ("body", "target", "notes"):
        assert f'href="#field-{field}"' in r.text, field
    assert "2,001 characters" in r.text
    r = post(user_client, f"/write/{doc_id}/suggest", {"goal": "spelling"})
    assert 'href="#field-body"' in r.text and 'href="#field-goal"' not in r.text
    assert fake_ai.calls == []
    page = user_client.get(f"/write/{doc_id}")
    assert "above" not in page.text and 'href="#field-target"' in page.text


def test_a_field_sent_twice_counts_once(user_client):
    r = post(user_client, "/write/new", {"title": ["First title", "Second title"], "kind": ["story", "article"]},
             follow_redirects=False)
    assert r.status_code == 303
    with db.session() as conn:
        row = conn.execute("SELECT title, kind FROM write_docs").fetchone()
    assert (row["title"], row["kind"]) == ("First title", "story")


def test_save_normalises_old_mac_line_breaks(user_client):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "One\rTwo\r\nThree")
    assert body_of(doc_id) == "One\nTwo\nThree"


def test_docx_export_survives_control_characters(user_client):
    from docx import Document
    doc_id = create_doc(user_client, "Pasted")
    save(user_client, doc_id, "Line one\x0bline two\x0cPage two\x00 with a nul\x07 and a bell.", title="Pasted\x01")
    r = user_client.get(f"/write/{doc_id}/export.docx")
    assert r.status_code == 200
    texts = [p.text for p in Document(io.BytesIO(r.content)).paragraphs]
    assert texts == ["Pasted", "Line one", "line two", "Page two with a nul and a bell."]
    # The person's own text is not changed by the export
    assert body_of(doc_id).startswith("Line one\x0bline two")


def test_versions_are_grouped_by_day_and_numbered(user_client):
    doc_id = create_doc(user_client)
    for text in ("One.", "Two.", "Three."):
        save(user_client, doc_id, text)
    with db.session() as conn:
        conn.execute("UPDATE write_versions SET created_at = '2026-01-02T10:00:00+00:00' WHERE body = 'One.'")
    page = user_client.get(f"/write/{doc_id}/versions")
    assert "<h2>Saved on 2 January 2026 (1 version)</h2>" in page.text
    assert re.search(r"<h2>Saved on \d{1,2} [A-Z][a-z]+ \d{4} \(2 versions\)</h2>", page.text)
    assert page.text.index("(2 versions)") < page.text.index("2 January 2026 (1 version)")
    assert not re.search(r"\d{4}-\d{2}-\d{2}", page.text.split("<main", 1)[1])
    assert re.search(r">Version 3: saved by you, \d\d:\d\d UTC \(current text\)</a>", page.text)
    assert ">Version 1: saved by you, 10:00 UTC</a>" in page.text
    assert not check_html(page.text).errors
    versions = versions_of(doc_id)
    v2 = user_client.get(f"/write/{doc_id}/versions/{versions[1]['id']}")
    assert h1(v2.text) == "Version 2 of My essay" and "Two." in v2.text
    assert "3 versions saved" in user_client.get(f"/write/{doc_id}").text
    assert "3 versions" in user_client.get(f"/write/{doc_id}/delete").text


def test_all_suggestions_dropped_is_explained_honestly(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [{"original": "not in the text", "proposed": "x", "reason": ""}])
    page = user_client.get(f"/write/{doc_id}/suggestions/{batch_id}")
    assert h1(page.text) == "No suggestions for My essay"
    assert "None of its suggestions could be used" in page.text
    assert "1 of the AI's 1 suggestion was left out" in page.text
    assert "The AI made 0" not in page.text and "more suggestion" not in page.text


# ----- ease of use: found by blind members using screen readers, braille and keyboards -------------

def test_ai_buttons_save_the_text_box_first(user_client, fake_ai):
    """The AI buttons sit in the same form as the text box, so edits are saved first and never thrown away."""
    doc_id = create_doc(user_client, target="Some Journal")
    page = user_client.get(f"/write/{doc_id}").text
    # One form: the AI buttons use formaction, so the text box goes with them
    form = page.split('action="/write/%d/save"' % doc_id, 1)[1].split("</form>", 1)[0]
    assert "<form" not in form
    for action, name in [("suggest", "Save and get suggestions"),
                         ("venues", "Save and search the web for places to publish"),
                         ("cover-letter", "Save and draft a cover letter for Some Journal")]:
        assert f'formaction="/write/{doc_id}/{action}">{name}</button>' in form, action
    assert 'name="body"' in form and ">Save my text</button>" in form
    # Never saved, the text box is filled in, and "Save and get suggestions" is chosen
    fake_ai.on("write.suggest", {"overall_note": "", "suggestions": [{"original": "teh", "proposed": "the", "reason": ""}]})
    r = post(user_client, f"/write/{doc_id}/suggest", {"body": BODY, "title": "My essay", "target": "Some Journal",
                                                       "goal": "spelling", "instructions": ""})
    assert r.status_code == 200 and "Saved your text as version 1, then asked the AI." in r.text
    assert body_of(doc_id) == BODY and [v["note"] for v in versions_of(doc_id)] == ["Saved by you"]
    assert "teh mat" in fake_ai.calls[-1][2]["prompt"]
    # An edit made just before asking for a search is what the AI reads
    fake_ai.on("write.venues", ResearchResult(text="Found one.", sources=[]))
    r = post(user_client, f"/write/{doc_id}/venues", {"body": BODY + "\n\nNew ending.", "title": "My essay",
                                                      "target": "Some Journal"})
    assert "Saved your text as version 2, then asked the AI." in r.text
    assert "New ending." in fake_ai.calls[-1][2]["prompt"]
    # Unchanged text is not saved again
    fake_ai.on("write.cover_letter", "Dear [Editor's name],")
    post(user_client, f"/write/{doc_id}/cover-letter", {"body": BODY + "\n\nNew ending.", "title": "My essay",
                                                        "target": "Some Journal", "notes": ""})
    assert len(versions_of(doc_id)) == 2


def test_ai_button_keeps_the_words_when_something_needs_checking(user_client, fake_ai):
    doc_id = create_doc(user_client)
    long_text = "A long pasted essay. " * 600
    # Nothing saved yet, and the instructions hold a card number: the text is kept in the box, the AI never runs
    r = post(user_client, f"/write/{doc_id}/suggest", {"body": long_text, "title": "My essay", "target": "",
                                                       "goal": "clearer", "instructions": "card 4111 1111 1111 1111"})
    assert r.status_code == 400 and h1(r.text) == "Please check 1 answer: My essay"
    assert long_text.strip() in r.text and 'href="#field-instructions"' in r.text
    assert "Saved your text as version 1. The AI has not started yet." in r.text
    assert body_of(doc_id) == long_text  # the text itself was fine, so it is saved
    assert 'value="clearer" selected' in r.text
    # A problem in the text box itself: nothing is saved, and every word is shown again
    r = post(user_client, f"/write/{doc_id}/cover-letter", {"body": "My card 4111 1111 1111 1111.", "title": "",
                                                            "target": "Some Journal", "notes": "I teach braille."})
    assert r.status_code == 400
    assert "My card 4111 1111 1111 1111." in r.text and "I teach braille." in r.text
    assert 'href="#field-title"' in r.text and 'href="#field-body"' in r.text
    assert body_of(doc_id) == long_text and fake_ai.calls == []


def test_reload_after_a_form_goes_back_to_the_document(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [{"original": "teh", "proposed": "the", "reason": ""}])
    sid = suggestion_rows(batch_id)[0]["id"]
    vid = versions_of(doc_id)[0]["id"]
    for path, back in [("/save", f"/write/{doc_id}"), ("/suggest", f"/write/{doc_id}"), ("/venues", f"/write/{doc_id}"),
                       ("/cover-letter", f"/write/{doc_id}"),
                       (f"/suggestions/{batch_id}/{sid}/accept", f"/write/{doc_id}/suggestions/{batch_id}"),
                       (f"/suggestions/{batch_id}/{sid}/reject", f"/write/{doc_id}/suggestions/{batch_id}"),
                       (f"/suggestions/{batch_id}/reject-all", f"/write/{doc_id}/suggestions/{batch_id}"),
                       (f"/versions/{vid}/restore", f"/write/{doc_id}/versions/{vid}")]:
        r = user_client.get(f"/write/{doc_id}{path}", follow_redirects=False)
        assert r.status_code == 303 and r.headers["location"] == back, path
    assert suggestion_rows(batch_id)[0]["status"] == "pending" and body_of(doc_id) == BODY


def test_edit_form_order_and_save_lands_on_the_notice(user_client):
    doc_id = create_doc(user_client)
    page = user_client.get(f"/write/{doc_id}").text
    assert page.index('id="field-title"') < page.index('id="field-body"') < page.index('id="field-target"')
    r = post(user_client, f"/write/{doc_id}/save", {"body": "Text.", "title": "My essay", "target": ""},
             follow_redirects=False)
    # No fragment, so the page opens at "Saved as version 1", which gets the focus.
    assert r.headers["location"] == f"/write/{doc_id}"
    page = user_client.get(f"/write/{doc_id}").text
    assert re.search(r"Last saved by you on \d{1,2} [A-Z][a-z]+ \d{4}, \d\d:\d\d UTC\. "
                     r"Times are in UTC, the same as London time in winter\.", page)
    assert not re.search(r"\d{4}-\d{2}-\d{2}", page.split("<main", 1)[1])


def test_deciding_moves_on_to_the_next_suggestion(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "Alpha one. Beta two. Gamma three.")
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [
        {"original": "Alpha one.", "proposed": "Alpha 1.", "reason": "Digits."},
        {"original": "Beta two.", "proposed": "Beta 2.", "reason": "Digits."},
        {"original": "Gamma three.", "proposed": "Gamma 3.", "reason": "Digits."},
    ])
    rows = suggestion_rows(batch_id)
    base = f"/write/{doc_id}/suggestions/{batch_id}"
    r = post(user_client, f"{base}/{rows[0]['id']}/accept", follow_redirects=False)
    assert r.headers["location"] == f"{base}#suggestion-2"
    page = user_client.get(base).text
    # The outcome is said at the top of the next suggestion, once, and the h1 says it too
    section = page.split('id="suggestion-2"', 1)[1].split("</section>", 1)[0]
    assert section.index("Suggestion 1 accepted. Your text has the change and is saved as version 2.") < section.index("<h2")
    assert page.count("Suggestion 1 accepted.") == 1
    assert h1(page) == "3 suggestions for My essay: suggestion 1 accepted"
    assert "<strong>Before:</strong> Alpha one." in page and "<strong>After (now in your text):</strong> Alpha 1." in page
    assert "2 waiting for your decision and 1 accepted." in page
    r = post(user_client, f"{base}/{rows[2]['id']}/reject", follow_redirects=False)
    assert r.headers["location"] == f"{base}#suggestion-2"
    page = user_client.get(base).text
    assert "<strong>Your words (kept):</strong> Gamma three." in page
    assert "<strong>Suggested, not used:</strong> Gamma 3." in page
    r = post(user_client, f"{base}/{rows[1]['id']}/reject", follow_redirects=False)
    assert r.headers["location"] == f"{base}#progress"  # none left: the counts
    page = user_client.get(base).text
    progress = page.split('id="progress"', 1)[1].split("</div>", 1)[0]
    assert "Suggestion 2 rejected. Your text is unchanged." in progress and "1 accepted and 2 rejected." in progress


def test_accept_all_is_last_and_asks_first(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "teh one, teh two.")
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [
        {"original": "teh one", "proposed": "the one", "reason": "Spelling."},
        {"original": "teh two", "proposed": "the two", "reason": "Spelling."},
    ])
    base = f"/write/{doc_id}/suggestions/{batch_id}"
    page = user_client.get(base).text
    # The first button on the page is for suggestion 1, not for all of them
    first_button = re.search(r"<button[^>]*>([^<]*)</button>", page.split("<main", 1)[1]).group(1)
    assert first_button == "Accept suggestion 1"
    assert f'<a href="{base}/accept-all">Accept all 2 waiting suggestions</a>' in page
    confirm = user_client.get(f"{base}/accept-all")
    assert confirm.status_code == 200 and h1(confirm.text) == "Accept all 2 waiting suggestions for My essay?"
    assert ">Accept every suggestion that still fits my text</button>" in confirm.text
    assert "No, go back and decide one at a time" in confirm.text
    assert not check_html(confirm.text).errors
    assert body_of(doc_id) == "teh one, teh two."  # opening the question changes nothing
    r = post(user_client, f"{base}/accept-all", follow_redirects=False)
    assert r.headers["location"] == f"{base}#progress"
    assert "Accepted 2 suggestions and saved your text as version 2." in user_client.get(base).text
    assert h1(user_client.get(f"{base}/accept-all").text) == "No suggestions are waiting for My essay"


def test_out_of_date_suggestion_is_a_notice_and_can_be_dismissed(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [{"original": "teh", "proposed": "the", "reason": ""},
                                                              {"original": "It was happy.", "proposed": "Glad.", "reason": ""}])
    save(user_client, doc_id, BODY.replace("teh", "thee").replace("It was happy.", "It was happy. It was happy."))
    base = f"/write/{doc_id}/suggestions/{batch_id}"
    page = user_client.get(base).text
    assert "Accept suggestion 1" not in page and "Dismiss suggestion 1 (no longer matches your text)" in page
    assert "Quoted from your earlier text:" in page and 'class="field-error"' not in page
    rows = suggestion_rows(batch_id)
    post(user_client, f"{base}/{rows[0]['id']}/reject", {"dismiss": "yes"})
    assert suggestion_rows(batch_id)[0]["status"] == "stale"
    # Accepting one that no longer fits is information, not an error alert
    r = post(user_client, f"{base}/{rows[1]['id']}/accept")
    assert 'role="status">Suggestion 2 no longer matches your text' in r.text and 'role="alert"' not in r.text
    assert suggestion_rows(batch_id)[1]["status"] == "stale"


def test_version_page_shows_what_is_different(user_client):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "# Start\n\nFirst paragraph.\n\nSecond paragraph.")
    save(user_client, doc_id, "# Start\n\nFirst paragraph, changed.\n\nSecond paragraph.\n\nA new ending.")
    first = versions_of(doc_id)[0]
    page = user_client.get(f"/write/{doc_id}/versions/{first['id']}").text
    assert ">What is different from your current text</h2>" in page and "Computed by the site, not the AI." in page
    assert "<strong>This version has:</strong> First paragraph." in page
    assert "<strong>Your current text has:</strong> First paragraph, changed." in page
    assert "<strong>Only your current text has:</strong> A new ending." in page
    assert "<h3>Start</h3>" in page and "<pre>" not in page  # real headings and paragraphs
    assert page.count(">Restore version 1</button>") == 2  # before and after the text
    assert not check_html(page).errors
    assert user_client.get(f"/write/{doc_id}/versions").text.count("(current text)</a>") == 1


EXPORT_LEVELS = "# Title again\n\n## Section\n\nText.\n\n#### Part\n\nMore text."


def test_exports_close_heading_gaps(user_client):
    from docx import Document
    doc_id = create_doc(user_client, "Title again")
    save(user_client, doc_id, EXPORT_LEVELS, title="Title again")
    document = Document(io.BytesIO(user_client.get(f"/write/{doc_id}/export.docx").content))
    headings = [(p.style.name, p.text) for p in document.paragraphs if p.style.name.startswith("Heading")]
    assert headings == [("Heading 1", "Title again"), ("Heading 2", "Section"), ("Heading 3", "Part")]
    page = user_client.get(f"/write/{doc_id}/export.html").text
    assert "<h2>Section</h2>" in page and "<h3>Part</h3>" in page and "<h4>" not in page
    tex = user_client.get(f"/write/{doc_id}/export.tex").text
    assert "\\section{Section}" in tex and "\\subsection{Part}" in tex
    txt = user_client.get(f"/write/{doc_id}/export.txt").text
    assert txt == "Title again\n\n1. Section\n\nText.\n\n1.1 Part\n\nMore text.\n" and "#" not in txt
    md = user_client.get(f"/write/{doc_id}/export.md").text
    assert md == "# Title again\n\n## Section\n\nText.\n\n### Part\n\nMore text.\n"
    export = user_client.get(f"/write/{doc_id}/export").text
    assert "File name: title-again.docx." in export and "Downloads folder" in export


def test_cover_letter_respects_disclosure_and_links_back(user_client, fake_ai):
    doc_id = create_doc(user_client, "Screen readers in class", "paper", "Journal of Examples")
    save(user_client, doc_id, BODY, title="Screen readers in class", target="Journal of Examples")
    fake_ai.on("write.cover_letter", "Dear [Editor's name],\n\n[Your name]")
    r = post(user_client, f"/write/{doc_id}/cover-letter", {"notes": ""}, follow_redirects=False)
    new_id = int(user_client.get(r.headers["location"], follow_redirects=False).headers["location"].rsplit("/", 1)[1])
    kind, purpose, kwargs = fake_ai.calls[-1]
    assert "Never mention that the writer is blind or has low vision" in kwargs["prompt"]
    letter = user_client.get(f"/write/{new_id}").text
    assert "This letter is a new document." in letter
    assert f'Your text "<a href="/write/{doc_id}">Screen readers in class</a>" is unchanged.' in letter
    assert 'id="publish-title"' not in letter and 'id="cover-title"' not in letter
    assert "Fill in everything in square brackets" in letter and "carefully" not in letter
    paper = user_client.get(f"/write/{doc_id}").text
    assert ">Cover letters drafted</h3>" in paper
    assert f'<a href="/write/{new_id}">Cover letter: Screen readers in class</a>' in paper


def test_index_table_has_a_caption_and_row_headers(user_client):
    doc_id = create_doc(user_client)
    page = user_client.get("/write").text
    assert "<caption>Your documents, most recently changed first</caption>" in page
    assert f'<th scope="row"><a href="/write/{doc_id}">My essay</a></th>' in page
    assert not re.search(r"\d{4}-\d{2}-\d{2}", page.split("<main", 1)[1])


def test_prompts_carry_only_task_rules():
    assert "screen reader" not in write_mod.SUGGEST_SYSTEM and "blind or has low vision" not in write_mod.SUGGEST_SYSTEM
    assert "never suggest euphemisms" in write_mod.SUGGEST_SYSTEM
    for system in (write_mod.SUGGEST_SYSTEM, write_mod.VENUES_SYSTEM, write_mod.COVER_LETTER_SYSTEM):
        assert "Put the summary first" not in system and "emoji" not in system


# ----- final walk-through fixes ----------------------------------------------------------------

def test_save_keeps_what_was_typed_for_the_ai(user_client):
    doc_id = create_doc(user_client)
    r = post(user_client, f"/write/{doc_id}/save", {"body": "Text.", "title": "My essay", "target": "",
                                                    "goal": "clearer", "instructions": "Keep British spelling.",
                                                    "notes": "I teach braille."}, follow_redirects=False)
    assert r.headers["location"] == f"/write/{doc_id}"
    page = user_client.get(f"/write/{doc_id}").text
    assert ">Keep British spelling.</textarea>" in page and ">I teach braille.</textarea>" in page
    assert '<option value="clearer" selected' in page
    # Emptied boxes are emptied for good, and anything that looks like a card number is never kept.
    post(user_client, f"/write/{doc_id}/save", {"body": "Text.", "title": "My essay", "target": "", "goal": "spelling",
                                                "instructions": "", "notes": "4111 1111 1111 1111"})
    page = user_client.get(f"/write/{doc_id}").text
    assert "4111" not in page and ">Keep British spelling.</textarea>" not in page
    with db.session() as conn:
        assert conn.execute("SELECT COUNT(*) FROM write_ai_fields WHERE doc_id = ?", (doc_id,)).fetchone()[0] == 0


def test_document_form_carries_the_saved_values_for_the_leave_warning(user_client):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, 'Saved "words".')
    # The form comes back with a problem: the box shows the new words, the form still says what is saved.
    r = post(user_client, f"/write/{doc_id}/save", {"body": "Unsaved words.", "title": "", "target": ""})
    assert r.status_code == 400 and ">Unsaved words.</textarea>" in r.text
    form = r.text.split('<form method="post" action="/write/', 1)[1].split(">", 1)[0]
    assert "data-leave-guard" in form and 'data-saved-body="Saved &#34;words&#34;."' in form
    assert 'data-saved-title="My essay"' in form and 'data-saved-instructions=""' in form


def test_new_document_notice_gives_only_the_next_step(user_client):
    r = post(user_client, "/write/new", {"title": "Letter", "kind": "article", "target": ""})
    assert 'Write or paste your text in "Your text", then choose "Save my text".' in html.unescape(r.text)
    assert "is ready" not in r.text
    r = post(user_client, "/write/new", {"title": "", "kind": "article", "target": ""})
    assert r.text.count("You can change it later.") == 1  # the hint only, not repeated by the error
    assert "Enter a title for your document." in r.text


def test_earlier_suggestions_are_numbered_and_lead_with_the_goal(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    first = ask_suggestions(user_client, doc_id, fake_ai, [])
    second = ask_suggestions(user_client, doc_id, fake_ai, [])
    page = user_client.get(f"/write/{doc_id}").text
    assert re.search(rf'<a href="/write/{doc_id}/suggestions/{second}">Set 2: Fix spelling and grammar only, '
                     r"\d{1,2} [A-Z][a-z]+ \d{4}, \d\d:\d\d UTC</a>", page)
    assert f'<a href="/write/{doc_id}/suggestions/{first}">Set 1: Fix spelling and grammar only, ' in page
    assert "Times are in UTC, the same as London time in winter." in page
    assert "(UTC, the same as London time in winter)" in user_client.get(f"/write/{doc_id}/suggestions/{first}").text


def test_accept_all_lists_what_would_change(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "teh one, teh two.")
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [
        {"original": "teh one", "proposed": "the one", "reason": "Spelling."},
        {"original": "teh two", "proposed": "the two", "reason": "Spelling."},
    ])
    save(user_client, doc_id, "teh one, two.")  # suggestion 2 no longer fits
    page = html.unescape(user_client.get(f"/write/{doc_id}/suggestions/{batch_id}/accept-all").text)
    assert '<li>Suggestion 1 replaces "teh" with "the".</li>' in page
    assert ('<li>Suggestion 2 replaces "teh" with "the". It no longer matches your text, so it would be left '
            'out.</li>') in page


def test_last_decision_offers_the_way_back(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    batch_id = ask_suggestions(user_client, doc_id, fake_ai, [
        {"original": "teh mat", "proposed": "the mat", "reason": "Spelling."}])
    base = f"/write/{doc_id}/suggestions/{batch_id}"
    sid = suggestion_rows(batch_id)[0]["id"]
    post(user_client, f"{base}/{sid}/reject")
    page = user_client.get(base).text
    progress = page.split('id="progress"', 1)[1].split("</div>", 1)[0]
    assert f'Every suggestion is decided. <a href="/write/{doc_id}">Back to My essay</a>.' in progress
    assert page.count(f'<a href="/write/{doc_id}">Back to My essay</a>') == 1


def test_version_differences_say_what_changed_and_where(user_client):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, "# Background\n\nMost apps assume you can see. Readers read teh words.")
    save(user_client, doc_id, "# Background\n\nMost apps assume that you can see. Readers read the words.")
    vid = versions_of(doc_id)[0]["id"]
    page = html.unescape(user_client.get(f"/write/{doc_id}/versions/{vid}").text)
    assert ('<strong>Where:</strong> In the section "Background", paragraph 1 of your current text.' in page)
    assert ('<strong>Changed since this version:</strong> adds "that" after "assume"; replaces "teh" with "the".'
            in page)
    assert "Times are in UTC, the same as London time in winter." in page
    only_old = write_mod.version_differences("One.\n\nTwo.", "One.")
    assert only_old == [{"version": "Two.", "current": "", "place": "Paragraph 2 of this version", "change": ""}]


def test_plain_text_export_keeps_numbers_the_member_wrote():
    doc = {"title": "T", "body": "# 1. Intro\n\nText.\n\n# 2. End\n\nMore."}
    assert write_mod.export_text(doc) == "T\n\n1. Intro\n\nText.\n\n2. End\n\nMore.\n"


def test_places_to_publish_name_the_site_of_each_source(user_client, fake_ai):
    doc_id = create_doc(user_client)
    save(user_client, doc_id, BODY)
    with db.session() as conn:
        cur = conn.execute("INSERT INTO write_research (doc_id, kind, content_md, sources_json, created_at)"
                           " VALUES (?, 'venues', 'Try the journal.', ?, ?)",
                           (doc_id, '[{"url": "https://www.example.org/submit", "title": "Submission guidelines"}]',
                            db.now()))
        rid = cur.lastrowid
    page = user_client.get(f"/write/{doc_id}/publish/{rid}").text
    assert 'rel="noopener noreferrer">Submission guidelines</a> (example.org)' in page
