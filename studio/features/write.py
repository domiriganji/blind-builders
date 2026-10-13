"""Write and publish: documents, AI suggestions shown as before and after, versions, export,
finding where to publish, and drafting a cover letter.

The AI never changes a person's words by itself. Every suggestion quotes an exact excerpt of the
text; the server keeps only suggestions whose excerpt really is in the text, and the person
accepts or rejects each one. Every change to the text is saved as a version.
"""
from __future__ import annotations

import io
import json
import re
import time
import unicodedata
from datetime import datetime, timezone
from difflib import SequenceMatcher
from urllib.parse import urlparse

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import Response

from .. import db, jobs, sensitive
from ..ai import AIError, AINotAvailable
from ..web import (checked_form, current_user, date_words, error_title, flash, get_ai, markdown, nest_headings,
                   redirect, render, require_user)

router = APIRouter(prefix="/write")
NAV = {"label": 'Write and publish', "url": '/write', "description": 'Write articles, papers and letters. The AI suggests changes one at a time and never silently rewords you. Get help finding a journal and submitting.'}

SCHEMA = """
CREATE TABLE IF NOT EXISTS write_docs (
    id INTEGER PRIMARY KEY,
    user_id INTEGER NOT NULL REFERENCES users(id),
    title TEXT NOT NULL,
    kind TEXT NOT NULL CHECK (kind IN ('article','paper','letter','story','other')),
    body TEXT NOT NULL DEFAULT '',
    target TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_write_docs_user ON write_docs(user_id);
CREATE TABLE IF NOT EXISTS write_versions (
    id INTEGER PRIMARY KEY,
    doc_id INTEGER NOT NULL REFERENCES write_docs(id) ON DELETE CASCADE,
    body TEXT NOT NULL,
    note TEXT NOT NULL DEFAULT '',
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_write_versions_doc ON write_versions(doc_id);
CREATE TABLE IF NOT EXISTS write_batches (
    id INTEGER PRIMARY KEY,
    doc_id INTEGER NOT NULL REFERENCES write_docs(id) ON DELETE CASCADE,
    goal TEXT NOT NULL,
    overall_note TEXT NOT NULL DEFAULT '',
    dropped_count INTEGER NOT NULL DEFAULT 0,
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_write_batches_doc ON write_batches(doc_id);
CREATE TABLE IF NOT EXISTS write_suggestions (
    id INTEGER PRIMARY KEY,
    doc_id INTEGER NOT NULL REFERENCES write_docs(id) ON DELETE CASCADE,
    batch_id INTEGER NOT NULL REFERENCES write_batches(id) ON DELETE CASCADE,
    number INTEGER NOT NULL,
    original TEXT NOT NULL,
    proposed TEXT NOT NULL,
    reason TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL DEFAULT 'pending' CHECK (status IN ('pending','accepted','rejected','stale')),
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_write_suggestions_batch ON write_suggestions(batch_id);
CREATE TABLE IF NOT EXISTS write_research (
    id INTEGER PRIMARY KEY,
    doc_id INTEGER NOT NULL REFERENCES write_docs(id) ON DELETE CASCADE,
    kind TEXT NOT NULL,
    content_md TEXT NOT NULL,
    sources_json TEXT NOT NULL DEFAULT '[]',
    created_at TEXT NOT NULL
);
CREATE INDEX IF NOT EXISTS idx_write_research_doc ON write_research(doc_id);
"""

KINDS = [
    ("article", "Article"),
    ("paper", "Paper for a journal or conference"),
    ("letter", "Letter"),
    ("story", "Story"),
    ("other", "Something else"),
]
KIND_NAMES = dict(KINDS)

GOALS = [
    ("spelling", "Fix spelling and grammar only"),
    ("clearer", "Make it clearer and keep my voice"),
    ("structure", "Check the structure and argument"),
    ("shorter", "Make it shorter"),
    ("publication", "Prepare it for the publication I named"),
]
GOAL_NAMES = dict(GOALS)
GOAL_INSTRUCTIONS = {
    "spelling": "Fix spelling, grammar and punctuation mistakes only. Do not change word choice, style, tone or meaning.",
    "clearer": "Make sentences clearer and easier to follow. Keep the writer's own voice, word choices and meaning "
               "wherever you can. Change as little as possible.",
    "structure": "Check the structure and the argument: the order of points, missing steps, claims without support, "
                 "weak links between paragraphs. Suggest concrete edits to specific passages. Explain larger issues "
                 "in the overall note.",
    "shorter": "Make the text shorter. Cut repetition, filler and needless words. Keep every key point and the writer's voice.",
    "publication": "Prepare the text for the publication named below: suggest changes so it fits what that publication "
                   "usually expects in tone, length and structure. Only state a publication's rules if you are sure "
                   "of them; otherwise say in the overall note that the writer should check the author guidelines.",
}

MAX_TITLE = 200
MAX_TARGET = 200
MAX_BODY = 300_000
MAX_INSTRUCTIONS = 2000
MAX_AI_BODY = 150_000          # longest text we send for suggestions
VENUE_EXCERPT = 8000           # how much of the text we send when searching for venues
MAX_SUGGESTIONS = 60
MAX_PROPOSED = 20_000          # longest replacement we accept from the AI
MAX_REASON = 1000
MAX_NOTE = 10_000
GONE = "This document was deleted before the AI started, so there was nothing to work on."
GONE_DURING = "This document was deleted while the AI was working, so the result was not kept."
BAD_ANSWER = "The AI's answer came back incomplete, so nothing was saved and your text was not changed. Please try again."
JOB_FRESH_SECONDS = 15 * 60    # a job older than this that never finished is treated as abandoned
OUTCOME_KEY = "write_outcome"
OUTCOME_SECONDS = 300

STATUS_TEXT = {
    "pending": "Waiting for your decision.",
    "accepted": "Accepted. Your text has this change.",
    "rejected": "Rejected. Your text was not changed.",
    "stale": "Out of date. The words it would change are no longer in your text, so nothing was changed.",
}

# The shared member note (ai.member_voice: who the writer is, tone, plain words, summary first) is put before
# every system prompt by AIClient, so these add only what each task needs (docs/VOICE.md, 4.18).
SUGGEST_SYSTEM = (
    "You suggest edits to a writer's own text on Blind Creators Studio. The writer accepts or rejects each one, "
    "so every suggestion must quote the text exactly. Respect the writer's voice. Never change the words the writer "
    "uses about disability or about themselves, such as \"blind\" or \"disabled\", and never suggest euphemisms "
    "for them."
)
VENUES_SYSTEM = "You help a writer on Blind Creators Studio choose where to submit their work."
COVER_LETTER_SYSTEM = ("You draft cover letters for a writer on Blind Creators Studio. The writer edits the letter "
                       "and sends it themselves.")
DRAFTED_NOTE = "Drafted by the AI"
COVER_PREFIX = "Cover letter: "

SUGGEST_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "required": ["overall_note", "suggestions"],
    "properties": {
        "overall_note": {"type": "string"},
        "suggestions": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["original", "proposed", "reason"],
                "properties": {
                    "original": {"type": "string"},
                    "proposed": {"type": "string"},
                    "reason": {"type": "string"},
                },
            },
        },
    },
}

# Looks like a payment card number: 15 or 16 digits, optionally in groups split by spaces or dashes.
_CARD_RE = re.compile(r"(?<!\d)(?:\d[ -]?){14,15}\d(?!\d)")
_LANG_RE = re.compile(r"^[A-Za-z]{2,3}(?:-[A-Za-z0-9]{2,8})?$")


# ----- small helpers ---------------------------------------------------------------------------

def when(iso: str | None) -> str:
    """'2026-10-12T14:03:22+00:00' -> '12 October 2026, 14:03 UTC'."""
    return date_words(iso, with_time=True) if iso else ""


def on_at(iso: str | None) -> str:
    """'2026-10-12T14:03:22+00:00' -> 'on 12 October 2026 at 14:03 UTC', for use inside a sentence."""
    if not iso:
        return ""
    day = date_words(iso)
    clock = iso[11:16] if len(iso) >= 16 and iso[13:14] == ":" else ""
    return f"on {day} at {clock} UTC" if clock else f"on {day}"


def _lower_first(text: str) -> str:
    return text[:1].lower() + text[1:] if text[:2] != "AI" else text


def saved_sentence(note: str, iso: str | None, prefix: str = "Last saved") -> str:
    """'Last saved by you on 13 October 2026 at 12:29 UTC.' or 'Last saved on ...: accepted suggestion 1.'"""
    if note == "Saved by you":
        return f"{prefix} by you {on_at(iso)}."
    if note == DRAFTED_NOTE:
        return f"Drafted by the AI {on_at(iso)}."
    return f"{prefix} {on_at(iso)}: {_lower_first(note)}." if note else f"{prefix} {on_at(iso)}."


def word_count(text: str) -> int:
    return len((text or "").split())


def plural(n: int, one: str, many: str | None = None) -> str:
    return f"{n} {one if n == 1 else (many or one + 's')}"


def _luhn_ok(digits: str) -> bool:
    total = 0
    for i, ch in enumerate(reversed(digits)):
        d = int(ch)
        if i % 2 == 1:
            d *= 2
            if d > 9:
                d -= 9
        total += d
    return total % 10 == 0


def looks_like_card_number(text: str) -> bool:
    for match in _CARD_RE.finditer(text or ""):
        digits = re.sub(r"\D", "", match.group(0))
        if len(digits) in (15, 16) and _luhn_ok(digits):
            return True
    return False


def _secret_error(text: str) -> str | None:
    """A plain message if the text looks like it holds a password, PIN, one-time code, API key or card number.
    Everything in a document can be sent to the AI, so such text is refused (docs/DEVELOPING.md)."""
    return sensitive.find_secret(text or "")


def _stored_secret_errors(doc: dict) -> dict[str, str]:
    """Problems in a document saved before these checks existed, so the AI is never sent them."""
    errors = {}
    for field in ("body", "title", "target"):
        problem = _secret_error(doc[field])
        if problem:
            errors[field] = problem + " Take it out and save, then ask again."
    return errors


def _slug(title: str) -> str:
    ascii_title = unicodedata.normalize("NFKD", title).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^A-Za-z0-9]+", "-", ascii_title).strip("-").lower()[:60].strip("-")
    return slug or "document"


def _text(form: dict, key: str) -> str:
    """One text value from a form. A field sent twice (for example by a double submission) counts once."""
    value = form.get(key, "")
    if isinstance(value, list):
        value = value[0] if value else ""
    return value if isinstance(value, str) else ""


_DOCX_BREAKS = re.compile("[\x0b\x0c\x1c\x1d\x1e\x85\u2028\u2029]")
_XML_ILLEGAL = re.compile("[\x00-\x08\x0e-\x1f\x7f\ufffe\uffff]")


def _xml_safe(text: str) -> str:
    """Text a Word file can hold. Line and page breaks pasted from other programs become line breaks;
    other invisible control characters, which Word files cannot contain, are left out."""
    return _XML_ILLEGAL.sub("", _DOCX_BREAKS.sub("\n", text or ""))


def _say(request: Request, url: str, text: str, note: str = "", kind: str = "success", place: int = 0) -> None:
    """Remember what just happened, so the next page's h1 (the first thing a screen reader reads) says it.
    `note` is a fuller sentence shown where the page opens (`place`: a suggestion number, or 0 for the top)."""
    request.session[OUTCOME_KEY] = {"url": url, "text": text[:300], "at": int(time.time()),
                                    "note": note[:600], "kind": kind, "place": place}


def _heard_all(request: Request, url: str) -> dict:
    """The outcome _say saved for this page, once. Empty if there is none."""
    outcome = request.session.pop(OUTCOME_KEY, None)
    if (isinstance(outcome, dict) and outcome.get("url") == url
            and time.time() - float(outcome.get("at") or 0) < OUTCOME_SECONDS):
        return outcome
    return {}


def _heard(request: Request, url: str) -> str:
    return str(_heard_all(request, url).get("text") or "")


def _begin_write(conn) -> None:
    """Take the write lock before reading, so two quick presses cannot both act on the same old text."""
    conn.execute("BEGIN IMMEDIATE")


def _running_job(user_id: int, kind: str) -> str | None:
    """A job of this kind that is still working, so pressing a button twice does not pay for the AI twice."""
    with db.session() as conn:
        row = conn.execute("SELECT id, created_at FROM jobs WHERE user_id = ? AND kind = ? AND status IN ('queued', 'running')"
                           " ORDER BY created_at DESC LIMIT 1", (user_id, kind)).fetchone()
    if row is None:
        return None
    try:
        age = (datetime.now(timezone.utc) - datetime.fromisoformat(row["created_at"])).total_seconds()
    except (TypeError, ValueError):
        return None
    return row["id"] if age < JOB_FRESH_SECONDS else None


def _start_job(request: Request, user: dict, doc: dict, kind: str, title: str, fn, *args, saved: str = ""):
    """Start AI work for a document, or send the person to the same work if it is already running.
    `saved` says the text was saved first ("Saved your text as version 3"), so the next page can say so."""
    kind = f"{kind}:{doc['id']}"
    running = _running_job(user["id"], kind)
    if running:
        if saved:
            flash(request, f"{saved}. The AI is already working on your earlier text, so nothing new was started and "
                           "you are not charged twice. When it is done, you can ask again.", "info")
        else:
            flash(request, "The AI is already working on this. Nothing new was started, so you are not charged twice.",
                  "info")
        return redirect(jobs.job_url(running))
    try:
        ai = get_ai(request, user)
    except AINotAvailable as exc:
        if saved:
            flash(request, f"{saved}.", "success")
        flash(request, exc.user_message, "error")
        _say(request, f"/write/{doc['id']}", f"{doc['title']}: the AI is not available")
        return redirect(f"/write/{doc['id']}")
    if saved:
        flash(request, f"{saved}, then asked the AI.", "success")
    return redirect(jobs.job_url(jobs.submit(user["id"], kind, title, fn, ai, *args, back_url=f"/write/{doc['id']}",
                                             item_ref=_item_ref(doc))))


def _item_ref(doc: dict) -> str:
    """What the document's jobs are about, so deleting the document removes them (see jobs.forget_item)."""
    return f"write:{doc['id']}:{doc['created_at']}"


def _not_found(what: str = "document") -> HTTPException:
    return HTTPException(status_code=404, detail=f"We could not find that {what}.")


def _clean_sources(sources) -> list[dict]:
    out = []
    for s in sources or []:
        url = getattr(s, "url", None) if not isinstance(s, dict) else s.get("url")
        title = getattr(s, "title", "") if not isinstance(s, dict) else s.get("title", "")
        quote = getattr(s, "quote", "") if not isinstance(s, dict) else s.get("quote", "")
        if not url or urlparse(str(url)).scheme not in ("http", "https"):
            continue
        out.append({"url": str(url), "title": str(title or url), "quote": str(quote or "")[:300]})
    return out


# ----- data access -----------------------------------------------------------------------------

def _owned_doc(doc_id: int, user: dict) -> dict:
    with db.session() as conn:
        row = conn.execute("SELECT * FROM write_docs WHERE id = ? AND user_id = ?", (doc_id, user["id"])).fetchone()
    if row is None:
        raise _not_found()
    return dict(row)


_JOB_DOC_SQL = "SELECT * FROM write_docs WHERE id = ? AND user_id = ? AND created_at = ?"


def _get_doc(conn, ctx, doc_id: int, created_at: str) -> dict | None:
    """The document a job was started for. Ids of deleted documents can be reused, so a job checks the owner
    and the creation time too and never touches a newer document that took the id."""
    row = conn.execute(_JOB_DOC_SQL, (doc_id, ctx.user_id, created_at)).fetchone()
    return dict(row) if row else None


def _job_doc(ctx, doc_id: int, created_at: str) -> dict | None:
    with db.session() as conn:
        return _get_doc(conn, ctx, doc_id, created_at)


def _versions(doc_id: int) -> list[dict]:
    """All versions, oldest first, each with its number (1 is the oldest)."""
    with db.session() as conn:
        rows = conn.execute("SELECT * FROM write_versions WHERE doc_id = ? ORDER BY id", (doc_id,)).fetchall()
    out = []
    for i, r in enumerate(rows, start=1):
        v = dict(r)
        v["number"] = i
        v["when"] = when(v["created_at"])
        v["words"] = word_count(v["body"])
        out.append(v)
    return out


def _version_summary(doc_id: int) -> tuple[int, dict | None]:
    """How many versions there are and the newest one, without loading every version's text."""
    with db.session() as conn:
        count = conn.execute("SELECT COUNT(*) AS n FROM write_versions WHERE doc_id = ?", (doc_id,)).fetchone()["n"]
        last = conn.execute("SELECT id, note, created_at FROM write_versions WHERE doc_id = ? ORDER BY id DESC LIMIT 1",
                            (doc_id,)).fetchone()
    if last is None:
        return count, None
    return count, dict(last, number=count, when=when(last["created_at"]))


def _versions_by_day(versions: list[dict]) -> list[tuple[str, list[dict]]]:
    """Versions grouped under the day they were saved, newest day first (versions should already be newest first)."""
    groups: list[tuple[str, list[dict]]] = []
    for v in versions:
        day = date_words((v["created_at"] or "")[:10]) or "an unknown date"
        if not groups or groups[-1][0] != day:
            groups.append((day, []))
        groups[-1][1].append(v)
    return groups


def _add_version(conn, doc_id: int, body: str, note: str) -> None:
    conn.execute("INSERT INTO write_versions (doc_id, body, note, created_at) VALUES (?, ?, ?, ?)",
                 (doc_id, body, note[:300], db.now()))


def _set_body(conn, doc_id: int, body: str, note: str) -> None:
    """Change the text and save it as a version. Every change to a document goes through here."""
    conn.execute("UPDATE write_docs SET body = ?, updated_at = ? WHERE id = ?", (body, db.now(), doc_id))
    _add_version(conn, doc_id, body, note)


def _owned_batch(doc: dict, batch_id: int) -> dict:
    with db.session() as conn:
        row = conn.execute("SELECT * FROM write_batches WHERE id = ? AND doc_id = ?", (batch_id, doc["id"])).fetchone()
    if row is None:
        raise _not_found("set of suggestions")
    return dict(row)


def _suggestions(batch_id: int) -> list[dict]:
    with db.session() as conn:
        rows = conn.execute("SELECT * FROM write_suggestions WHERE batch_id = ? ORDER BY number", (batch_id,)).fetchall()
    return [dict(r) for r in rows]


def _batches(doc_id: int) -> list[dict]:
    with db.session() as conn:
        rows = conn.execute(
            "SELECT b.*, (SELECT COUNT(*) FROM write_suggestions s WHERE s.batch_id = b.id) AS total,"
            " (SELECT COUNT(*) FROM write_suggestions s WHERE s.batch_id = b.id AND s.status = 'pending') AS waiting"
            " FROM write_batches b WHERE b.doc_id = ? ORDER BY b.id DESC", (doc_id,)).fetchall()
    out = []
    for r in rows:
        b = dict(r)
        b["when"] = when(b["created_at"])
        b["goal_name"] = GOAL_NAMES.get(b["goal"], b["goal"])
        out.append(b)
    return out


def _research_list(doc_id: int) -> list[dict]:
    with db.session() as conn:
        rows = conn.execute("SELECT id, kind, created_at FROM write_research WHERE doc_id = ? ORDER BY id DESC",
                            (doc_id,)).fetchall()
    return [dict(r, when=when(r["created_at"])) for r in rows]


def _doc_view(doc: dict) -> dict:
    view = dict(doc)
    view["words"] = word_count(doc["body"])
    view["word_text"] = plural(view["words"], "word")
    view["kind_name"] = KIND_NAMES.get(doc["kind"], doc["kind"])
    view["updated"] = when(doc["updated_at"])
    view["updated_day"] = date_words(doc["updated_at"])
    return view


def _first_note(conn, doc_id: int) -> str:
    row = conn.execute("SELECT note FROM write_versions WHERE doc_id = ? ORDER BY id LIMIT 1", (doc_id,)).fetchone()
    return row["note"] if row else ""


def _drafted_for(doc: dict) -> tuple[bool, dict | None]:
    """Whether this document is a cover letter the AI drafted, and if so the document it was drafted for
    (found by its title, so it is None if either was renamed or deleted)."""
    with db.session() as conn:
        if doc["kind"] != "letter" or _first_note(conn, doc["id"]) != DRAFTED_NOTE:
            return False, None
        if not doc["title"].startswith(COVER_PREFIX):
            return True, None
        row = conn.execute("SELECT id, title FROM write_docs WHERE user_id = ? AND title = ? AND id != ?"
                           " AND created_at <= ? ORDER BY id DESC LIMIT 1",
                           (doc["user_id"], doc["title"][len(COVER_PREFIX):], doc["id"], doc["created_at"])).fetchone()
    return True, dict(row) if row else None


def _letters_for(doc: dict) -> list[dict]:
    """Cover letters the AI drafted for this document, newest first."""
    with db.session() as conn:
        rows = conn.execute(
            "SELECT d.id, d.title, d.created_at FROM write_docs d WHERE d.user_id = ? AND d.kind = 'letter'"
            " AND d.title = ? AND d.id != ? AND (SELECT v.note FROM write_versions v WHERE v.doc_id = d.id"
            " ORDER BY v.id LIMIT 1) = ? ORDER BY d.id DESC",
            (doc["user_id"], (COVER_PREFIX + doc["title"])[:MAX_TITLE], doc["id"], DRAFTED_NOTE)).fetchall()
    return [dict(r, when=when(r["created_at"])) for r in rows]


# ----- pages: list and create ------------------------------------------------------------------

def _new_values(title: str = "", kind: str = "article", target: str = "") -> dict:
    return {"title": title[:MAX_TITLE], "kind": kind if kind in KIND_NAMES else "article", "target": target[:MAX_TARGET]}


@router.get("")
def index(request: Request):
    user = current_user(request)
    if user is None:
        return render(request, "write/index.html", title="Write and publish", docs=[], values=_new_values(),
                      errors={}, kinds=KINDS)
    with db.session() as conn:
        rows = conn.execute("SELECT * FROM write_docs WHERE user_id = ? ORDER BY updated_at DESC, id DESC",
                            (user["id"],)).fetchall()
    docs = [_doc_view(dict(r)) for r in rows]
    outcome = _heard(request, "/write")
    return render(request, "write/index.html", title=f"Write and publish: {outcome}" if outcome else "Write and publish",
                  docs=docs, values=_new_values(), errors={}, kinds=KINDS)


@router.get("/new")
def new_form(request: Request, title: str = "", kind: str = "article", target: str = "",
             user: dict = Depends(require_user)):
    return render(request, "write/new.html", title="Start a new document", values=_new_values(title, kind, target),
                  errors={}, kinds=KINDS)


@router.post("/new")
async def create(request: Request, user: dict = Depends(require_user)):
    form = await checked_form(request)
    values = {k: _text(form, k).strip() for k in ("title", "kind", "target")}
    errors: dict[str, str] = {}
    if not values["title"]:
        errors["title"] = "Enter a title for your document. You can change it later."
    elif len(values["title"]) > MAX_TITLE:
        errors["title"] = f"Use a title of {MAX_TITLE} characters or fewer."
    if values["kind"] not in KIND_NAMES:
        errors["kind"] = "Choose what kind of writing this is."
    if len(values["target"]) > MAX_TARGET:
        errors["target"] = f"Use {MAX_TARGET} characters or fewer for the publication's name."
    for field in ("title", "target"):
        if field not in errors and _secret_error(values[field]):
            errors[field] = _secret_error(values[field])
            values[field] = ""  # never show the secret again
    if errors:
        return render(request, "write/new.html", title=error_title("Start a new document", errors),
                      values=values, errors=errors, kinds=KINDS, status_code=400)
    now = db.now()
    with db.session() as conn:
        cur = conn.execute(
            "INSERT INTO write_docs (user_id, title, kind, body, target, created_at, updated_at) VALUES (?, ?, ?, '', ?, ?, ?)",
            (user["id"], values["title"], values["kind"], values["target"], now, now))
        doc_id = cur.lastrowid
    flash(request, f"Your document \"{values['title']}\" is ready. Write or paste your text, then save it.", "success")
    _say(request, f"/write/{doc_id}", f"{values['title']}: new document started")
    return redirect(f"/write/{doc_id}")


# ----- the document page -----------------------------------------------------------------------

# The document page is one form: "Save my text" and each AI button send the text box with them, so the AI
# always works on what the member sees, and nothing typed is lost when they choose an AI button.
AI_FIELDS = ("goal", "instructions", "notes")


def _render_doc(request: Request, doc: dict, *, title: str | None = None, values: dict | None = None,
                errors: dict | None = None, status_code: int = 200):
    version_count, last_version = _version_summary(doc["id"])
    vals = {"title": doc["title"], "body": doc["body"], "target": doc["target"], "goal": "spelling",
            "instructions": "", "notes": ""}
    vals.update(values or {})
    is_letter, drafted_for = _drafted_for(doc)
    saved_line = (saved_sentence(last_version["note"], last_version["created_at"]) if last_version
                  else "Nothing saved yet.")
    return render(request, "write/doc.html", title=title or doc["title"], doc=_doc_view(doc), values=vals,
                  errors=errors or {}, goals=GOALS, version_count=version_count, last_version=last_version,
                  saved_line=saved_line, is_cover_letter=is_letter, drafted_for=drafted_for,
                  letters=[] if is_letter else _letters_for(doc),
                  batches=_batches(doc["id"]), research=_research_list(doc["id"]), status_code=status_code)


def _fix_problems(request: Request, doc: dict, values: dict, errors: dict):
    return _render_doc(request, doc, title=error_title(doc["title"], errors), values=values, errors=errors,
                       status_code=400)


@router.get("/{doc_id}")
def doc_page(request: Request, doc_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    outcome = _heard(request, f"/write/{doc_id}")
    return _render_doc(request, doc, title=outcome or None)


@router.get("/{doc_id}/save")
@router.get("/{doc_id}/suggest")
@router.get("/{doc_id}/venues")
@router.get("/{doc_id}/cover-letter")
def back_to_document(doc_id: int, user: dict = Depends(require_user)):
    """A reload or a restored tab after a form came back with a question lands here: go to the document."""
    _owned_doc(doc_id, user)
    return redirect(f"/write/{doc_id}")


def _edit_values(form: dict) -> dict:
    return {"body": _text(form, "body").replace("\r\n", "\n").replace("\r", "\n"),
            "title": _text(form, "title").strip(), "target": _text(form, "target").strip()}


def _kept_values(form: dict) -> dict:
    """Everything the member typed on the document page, to show again if something needs checking."""
    kept = {k: _text(form, k).strip() for k in AI_FIELDS if k in form}
    if "body" in form:
        kept.update(_edit_values(form))
    return kept


def _edit_errors(values: dict) -> dict[str, str]:
    body, new_title, target = values["body"], values["title"], values["target"]
    errors: dict[str, str] = {}
    if len(body) > MAX_BODY:
        errors["body"] = (f"Your text has {len(body):,} characters, and one document can hold up to {MAX_BODY:,}. "
                          "Please split it into two documents.")
    elif _secret_error(body):
        errors["body"] = _secret_error(body) + " Then save again."
    if not new_title:
        errors["title"] = "Enter a title for your document."
    elif len(new_title) > MAX_TITLE:
        errors["title"] = f"Use a title of {MAX_TITLE} characters or fewer."
    if len(target) > MAX_TARGET:
        errors["target"] = f"Use {MAX_TARGET} characters or fewer for the publication's name."
    elif _secret_error(target):
        errors["target"] = _secret_error(target)
    if "title" not in errors and _secret_error(new_title):
        errors["title"] = _secret_error(new_title)
    return errors


def _store_edits(doc_id: int, values: dict) -> int:
    """Save the title, publication and text. Returns the new version's number, or 0 if the text had not changed."""
    with db.session() as conn:
        _begin_write(conn)
        # Read the text again under the lock: a second press of Save must not save a second identical version.
        current = conn.execute("SELECT body FROM write_docs WHERE id = ?", (doc_id,)).fetchone()["body"]
        conn.execute("UPDATE write_docs SET title = ?, target = ?, updated_at = ? WHERE id = ?",
                     (values["title"], values["target"], db.now(), doc_id))
        last = conn.execute("SELECT body FROM write_versions WHERE doc_id = ? ORDER BY id DESC LIMIT 1",
                            (doc_id,)).fetchone()
        if values["body"] == current and not (last is None and values["body"] != ""):
            return 0
        _set_body(conn, doc_id, values["body"], "Saved by you")
        return conn.execute("SELECT COUNT(*) AS n FROM write_versions WHERE doc_id = ?", (doc_id,)).fetchone()["n"]


async def _save_first(request: Request, doc_id: int, user: dict) -> tuple[dict, dict, dict, str]:
    """For the AI buttons: check the form, then save the text box first if it came with the form and is fine,
    so the AI works on what the member sees. Returns (the document as saved, what the member typed, problems
    with the text box, and "Saved your text as version 3" or "")."""
    form = await checked_form(request)
    doc = _owned_doc(doc_id, user)
    kept = _kept_values(form)
    if "body" not in form:
        return doc, kept, {}, ""
    errors = _edit_errors(kept)
    if errors:
        return doc, kept, errors, ""
    version = _store_edits(doc_id, kept)
    return _owned_doc(doc_id, user), kept, {}, (f"Saved your text as version {version}" if version else "")


def _ai_problems(request: Request, doc: dict, kept: dict, errors: dict, saved: str):
    """Show the document again with what needs checking before the AI can start."""
    if saved:
        flash(request, f"{saved}. The AI has not started yet.", "success")
    return _fix_problems(request, doc, kept, errors)


@router.post("/{doc_id}/save")
async def save(request: Request, doc_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    form = await checked_form(request)
    kept = _kept_values(form)
    values = _edit_values(form)
    errors = _edit_errors(values)
    if errors:
        return _fix_problems(request, doc, kept, errors)
    version = _store_edits(doc_id, values)
    if version:
        flash(request, f"Saved as version {version}. {plural(word_count(values['body']), 'word')}.", "success")
        _say(request, f"/write/{doc_id}", f"{values['title']}: saved as version {version}")
    else:
        flash(request, "Saved. Your text had not changed, so no new version was needed.", "success")
        _say(request, f"/write/{doc_id}", f"{values['title']}: saved, the text had not changed")
    return redirect(f"/write/{doc_id}#edit-title")


# ----- suggestions -----------------------------------------------------------------------------

@router.post("/{doc_id}/suggest")
async def suggest(request: Request, doc_id: int, user: dict = Depends(require_user)):
    doc, kept, errors, saved = await _save_first(request, doc_id, user)
    goal = kept.get("goal", "")
    instructions = kept.get("instructions", "")
    if not errors:  # the text box is saved, so check the saved document
        if not doc["body"].strip():
            errors["body"] = ("Your document has no text yet. Write or paste it in \"Your text\", "
                              "then choose \"Save and get suggestions\".")
        elif len(doc["body"]) > MAX_AI_BODY:
            errors["body"] = (f"Your text has {len(doc['body']):,} characters, and one round of suggestions can take up "
                              f"to {MAX_AI_BODY:,}. Split it into two documents and ask about each one.")
        if goal == "publication" and not doc["target"]:
            errors["target"] = ("Fill in the publication you are aiming for, so the AI knows what to prepare your text "
                                "for. Or choose another goal.")
        for field, problem in _stored_secret_errors(doc).items():
            errors.setdefault(field, problem)
    if goal not in GOAL_NAMES:
        errors["goal"] = "Choose what you want the AI to look for."
    if len(instructions) > MAX_INSTRUCTIONS:
        errors["instructions"] = (f"Your instructions are {len(instructions):,} characters long. "
                                  f"Shorten them to {MAX_INSTRUCTIONS:,} characters or fewer.")
    elif _secret_error(instructions):
        errors["instructions"] = _secret_error(instructions)
    if errors:
        return _ai_problems(request, doc, kept, errors, saved)
    return _start_job(request, user, doc, "write.suggest", f"Suggestions for {doc['title']}", _suggest_job, doc_id,
                      doc["created_at"], goal, instructions, saved=saved)


def filter_suggestions(body: str, raw: list[dict]) -> tuple[list[dict], int]:
    """Keep suggestions whose original is an exact excerpt of body, found exactly once (so the change can go
    in one place only), and that change something.
    A suggestion with a missing or wrongly typed original or proposed text is dropped, never guessed at:
    a missing proposed text must not turn into "remove these words".
    Returns (kept, in document order) and how many were dropped."""
    kept: list[tuple[int, int, dict]] = []
    dropped = 0
    seen: set[tuple[str, str]] = set()
    for i, s in enumerate(raw or []):
        s = s if isinstance(s, dict) else {}
        original, proposed, reason = s.get("original"), s.get("proposed"), s.get("reason")
        if not isinstance(original, str) or not isinstance(proposed, str) or len(proposed) > MAX_PROPOSED:
            dropped += 1
            continue
        reason = reason.strip()[:MAX_REASON] if isinstance(reason, str) else ""
        pos = body.find(original) if original else -1
        if pos < 0 or original == proposed or (original, proposed) in seen or body.count(original) > 1:
            dropped += 1
            continue
        seen.add((original, proposed))
        kept.append((pos, i, {"original": original, "proposed": proposed, "reason": reason}))
    kept.sort(key=lambda t: (t[0], t[1]))
    if len(kept) > MAX_SUGGESTIONS:
        dropped += len(kept) - MAX_SUGGESTIONS
        kept = kept[:MAX_SUGGESTIONS]
    return [k[2] for k in kept], dropped


def _suggest_job(ctx, ai, doc_id: int, created_at: str, goal: str, instructions: str) -> str:
    doc = _job_doc(ctx, doc_id, created_at)
    if doc is None:
        raise AIError(GONE)
    body = doc["body"]
    ctx.progress("The AI is reading your text")
    prompt_parts = [
        f"Goal: {GOAL_INSTRUCTIONS[goal]}",
        f"Kind of writing: {KIND_NAMES.get(doc['kind'], doc['kind'])}.",
        f"Title: {doc['title']}",
    ]
    if doc["target"]:
        prompt_parts.append(f"Publication the writer is aiming for: {doc['target']}")
    if instructions:
        prompt_parts.append(f"Extra instructions from the writer: {instructions}")
    prompt_parts.append(
        "Rules for each suggestion:\n"
        "- original: copy an exact excerpt of the text, character for character, including punctuation. "
        "Keep it short: one sentence or less. Pick an excerpt that appears only once if you can. Suggestions must not overlap.\n"
        "- proposed: the new wording that replaces exactly that excerpt.\n"
        "- reason: one short sentence in plain words saying why.\n"
        f"Give at most {MAX_SUGGESTIONS // 2} suggestions, the most useful first. If nothing needs changing, give none.\n"
        "overall_note: two to four plain sentences about the text as a whole, summary first. Do not repeat the suggestions."
    )
    prompt_parts.append(f"The text:\n<text>\n{body}\n</text>")
    result = ai.structured(system=SUGGEST_SYSTEM, prompt="\n\n".join(prompt_parts), schema=SUGGEST_SCHEMA,
                           purpose="write.suggest")
    ctx.progress("Checking each suggestion against your text")
    if not isinstance(result, dict) or not isinstance(result.get("suggestions"), list):
        # Showing "no suggestions" here would wrongly tell the writer their text needs nothing.
        raise AIError(BAD_ANSWER)
    kept, dropped = filter_suggestions(body, result["suggestions"])
    note = result.get("overall_note")
    note = note.strip()[:MAX_NOTE] if isinstance(note, str) else ""
    now = db.now()
    with db.session() as conn:
        _begin_write(conn)
        if _get_doc(conn, ctx, doc_id, created_at) is None:
            raise AIError(GONE_DURING)
        cur = conn.execute(
            "INSERT INTO write_batches (doc_id, goal, overall_note, dropped_count, created_at) VALUES (?, ?, ?, ?, ?)",
            (doc_id, goal, note, dropped, now))
        batch_id = cur.lastrowid
        for number, s in enumerate(kept, start=1):
            conn.execute(
                "INSERT INTO write_suggestions (doc_id, batch_id, number, original, proposed, reason, status, created_at)"
                " VALUES (?, ?, ?, ?, ?, ?, 'pending', ?)",
                (doc_id, batch_id, number, s["original"], s["proposed"], s["reason"], now))
    return f"/write/{doc_id}/suggestions/{batch_id}"


def _counts_sentence(counts: dict) -> str:
    """'2 waiting for your decision, 1 accepted.' Only the counts that are not zero."""
    parts = [f"{counts['pending']} waiting for your decision" if counts["pending"] else ""]
    parts += [f"{counts[k]} {word}" for k, word in (("accepted", "accepted"), ("rejected", "rejected"),
                                                    ("stale", "out of date")) if counts[k]]
    parts = [p for p in parts if p]
    if not parts:
        return ""
    text = ", ".join(parts[:-1]) + (" and " if len(parts) > 1 else "") + parts[-1]
    return text[:1].upper() + text[1:] + "."


@router.get("/{doc_id}/suggestions/{batch_id}")
def suggestions_page(request: Request, doc_id: int, batch_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    batch = _owned_batch(doc, batch_id)
    items = _suggestions(batch_id)
    for s in items:
        s["matches"] = doc["body"].count(s["original"]) == 1
        s["change"] = describe_change(s["original"], s["proposed"])
        s["place"] = where_in_text(doc["body"], s["original"]) if s["status"] == "pending" and s["matches"] else ""
    n = len(items)
    counts = {k: sum(1 for s in items if s["status"] == k) for k in ("pending", "accepted", "rejected", "stale")}
    next_pending = next((s for s in items if s["status"] == "pending"), None)
    if n == 0:
        title = f"No suggestions for {doc['title']}"
    else:
        title = f"{plural(n, 'suggestion')} for {doc['title']}"
    outcome = _heard_all(request, f"/write/{doc_id}/suggestions/{batch_id}")
    if outcome.get("text"):
        title = f"{title}: {outcome['text']}"
    note = ({"text": outcome["note"], "kind": outcome.get("kind") or "success", "place": int(outcome.get("place") or 0)}
            if outcome.get("note") else None)
    if note and note["place"] and not any(s["number"] == note["place"] and s["status"] == "pending" for s in items):
        note["place"] = 0
    return render(request, "write/suggestions.html", title=title, doc=_doc_view(doc),
                  batch=dict(batch, when=when(batch["created_at"]), goal_name=GOAL_NAMES.get(batch["goal"], batch["goal"])),
                  items=items, n=n, counts=counts, counts_sentence=_counts_sentence(counts), next_pending=next_pending,
                  status_text=STATUS_TEXT, note=note, all_waiting=_all_waiting(counts["pending"]))


@router.get("/{doc_id}/suggestions/{batch_id}/accept-all")
def accept_all_confirm(request: Request, doc_id: int, batch_id: int, user: dict = Depends(require_user)):
    """Accepting every suggestion rewrites a lot of text at once, so it is confirmed first."""
    doc = _owned_doc(doc_id, user)
    batch = _owned_batch(doc, batch_id)
    waiting = sum(1 for s in _suggestions(batch_id) if s["status"] == "pending")
    if waiting:
        title = f"Accept {_all_waiting(waiting)} for {doc['title']}?"
    else:
        title = f"No suggestions are waiting for {doc['title']}"
    return render(request, "write/accept_all.html", title=title, doc=_doc_view(doc), batch=batch, waiting=waiting)


def _all_waiting(count: int) -> str:
    """'all 3 waiting suggestions', or 'the 1 waiting suggestion'."""
    return f"all {count} waiting suggestions" if count != 1 else "the 1 waiting suggestion"


_QUOTE_MAX = 80


def _quote(words: list[str]) -> str:
    text = " ".join(words)
    return f'"{text[:_QUOTE_MAX - 3]}..."' if len(text) > _QUOTE_MAX else f'"{text}"'


def describe_change(original: str, proposed: str) -> str:
    """What a suggestion changes, word by word, worked out by the site (not the AI), so a short change can be
    found without comparing two long sentences: 'replaces "teh" with "the".'"""
    old, new = original.split(), proposed.split()
    if old == new:
        return "changes only spaces or line breaks."
    if not new:
        text = f"removes {_quote(old)}"
        return text if text.endswith(('."', '!"', '?"')) else text + "."
    ops = [op for op in SequenceMatcher(None, old, new, autojunk=False).get_opcodes() if op[0] != "equal"]
    if len(ops) > 3:
        return f"makes {len(ops)} changes in this passage."
    parts = []
    for tag, i1, i2, j1, j2 in ops:
        if tag == "replace":
            parts.append(f"replaces {_quote(old[i1:i2])} with {_quote(new[j1:j2])}")
        elif tag == "delete":
            parts.append(f"removes {_quote(old[i1:i2])}")
        elif i1 > 0:
            parts.append(f"adds {_quote(new[j1:j2])} after {_quote(old[i1 - 1:i1])}")
        elif old:
            parts.append(f"adds {_quote(new[j1:j2])} before {_quote(old[:1])}")
        else:
            parts.append(f"adds {_quote(new[j1:j2])}")
    text = "; ".join(parts)
    return text if text.endswith(('."', '!"', '?"')) else text + "."


_MD_HEADING = re.compile(r"^\s*#{1,6}\s+")


def where_in_text(body: str, excerpt: str, width: int = 60) -> str:
    """Where the excerpt is, place first, with a few words either side and no Markdown marks, so a short
    change like one word can be found: 'In the section "Methods", paragraph 2: ...we paid teh fee...'."""
    pos = body.find(excerpt) if excerpt else -1
    if pos < 0:
        return ""
    section, paragraph, in_paragraph, para_start = "", 0, False, 0
    offset = 0
    for line in body.split("\n"):
        end = offset + len(line)
        heading = _heading(line.strip())
        if heading:
            if offset <= pos <= end:
                return f'In the heading "{heading[1]}".'
            section, paragraph, in_paragraph = heading[1], 0, False
        elif not line.strip():
            in_paragraph = False
        elif not in_paragraph:
            paragraph, in_paragraph, para_start = paragraph + 1, True, offset
        if offset <= pos <= end:
            break
        offset = end + 1
    para_end = body.find("\n\n", pos + len(excerpt))
    para_end = len(body) if para_end < 0 else para_end
    start, stop = max(para_start, pos - width), min(para_end, pos + len(excerpt) + width)
    before, after = body[start:pos], body[pos + len(excerpt):stop]
    if start > para_start and " " in before:
        before = before[before.index(" ") + 1:]
    if stop < para_end and " " in after:
        after = after[:after.rindex(" ")]
    text = (("..." if start > para_start else "") + before + excerpt + after + ("..." if stop < para_end else ""))
    text = " ".join(_MD_HEADING.sub("", ln) for ln in text.split("\n")).split()
    place = (f'In the section "{section}", paragraph {paragraph}' if section else f"Paragraph {paragraph}")
    return f"{place}: {' '.join(text)}"


def _owned_suggestion(doc: dict, batch_id: int, suggestion_id: int) -> dict:
    with db.session() as conn:
        row = conn.execute("SELECT * FROM write_suggestions WHERE id = ? AND batch_id = ? AND doc_id = ?",
                           (suggestion_id, batch_id, doc["id"])).fetchone()
    if row is None:
        raise _not_found("suggestion")
    return dict(row)


def _next_place(conn, batch_id: int, after: int) -> int:
    """The next suggestion still waiting after this one (or the first waiting one), or 0 if none are left."""
    rows = [r["number"] for r in conn.execute(
        "SELECT number FROM write_suggestions WHERE batch_id = ? AND status = 'pending' ORDER BY number", (batch_id,))]
    return next((n for n in rows if n > after), rows[0] if rows else 0)


def _back_to(doc_id: int, batch_id: int, place: int) -> str:
    """The suggestions page, opened at the next suggestion to decide, or at the counts when none are left."""
    return f"/write/{doc_id}/suggestions/{batch_id}" + (f"#suggestion-{place}" if place else "#progress")


@router.get("/{doc_id}/suggestions/{batch_id}/{suggestion_id}/accept")
@router.get("/{doc_id}/suggestions/{batch_id}/{suggestion_id}/reject")
@router.get("/{doc_id}/suggestions/{batch_id}/reject-all")
def back_to_suggestions(doc_id: int, batch_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    _owned_batch(doc, batch_id)
    return redirect(f"/write/{doc_id}/suggestions/{batch_id}")


@router.post("/{doc_id}/suggestions/{batch_id}/{suggestion_id}/accept")
async def accept(request: Request, doc_id: int, batch_id: int, suggestion_id: int, user: dict = Depends(require_user)):
    await checked_form(request)
    doc = _owned_doc(doc_id, user)
    _owned_batch(doc, batch_id)
    s = _owned_suggestion(doc, batch_id, suggestion_id)
    here = f"/write/{doc_id}/suggestions/{batch_id}"
    kind = "success"
    with db.session() as conn:
        # Under the write lock, so pressing Accept twice cannot apply the change twice (to a second "teh", say).
        _begin_write(conn)
        status = conn.execute("SELECT status FROM write_suggestions WHERE id = ?", (s["id"],)).fetchone()["status"]
        body = conn.execute("SELECT body FROM write_docs WHERE id = ?", (doc_id,)).fetchone()["body"]
        if status != "pending":
            outcome = f"suggestion {s['number']} was already decided"
            message, kind = f"Suggestion {s['number']} was already decided: {STATUS_TEXT[status]}", "info"
        elif body.count(s["original"]) == 1:
            _set_body(conn, doc_id, body.replace(s["original"], s["proposed"], 1), f"Accepted suggestion {s['number']}")
            conn.execute("UPDATE write_suggestions SET status = 'accepted' WHERE id = ?", (s["id"],))
            version = conn.execute("SELECT COUNT(*) AS n FROM write_versions WHERE doc_id = ?", (doc_id,)).fetchone()["n"]
            outcome = f"suggestion {s['number']} accepted"
            message = f"Suggestion {s['number']} accepted. Your text has the change and is saved as version {version}."
        else:
            # The member edited their own text since. That is information, not an error (docs/VOICE.md, 4.12).
            conn.execute("UPDATE write_suggestions SET status = 'stale' WHERE id = ?", (s["id"],))
            outcome = f"suggestion {s['number']} is out of date, nothing changed"
            message, kind = (f"Suggestion {s['number']} no longer matches your text, because the words it would change "
                             "are not there any more, or are now there more than once. Nothing was changed.", "info")
        place = _next_place(conn, batch_id, s["number"])
    _say(request, here, outcome, message, kind, place)
    return redirect(_back_to(doc_id, batch_id, place))


@router.post("/{doc_id}/suggestions/{batch_id}/{suggestion_id}/reject")
async def reject(request: Request, doc_id: int, batch_id: int, suggestion_id: int, user: dict = Depends(require_user)):
    form = await checked_form(request)
    doc = _owned_doc(doc_id, user)
    _owned_batch(doc, batch_id)
    s = _owned_suggestion(doc, batch_id, suggestion_id)
    here = f"/write/{doc_id}/suggestions/{batch_id}"
    dismiss = _text(form, "dismiss") == "yes"
    with db.session() as conn:
        _begin_write(conn)
        body = conn.execute("SELECT body FROM write_docs WHERE id = ?", (doc_id,)).fetchone()["body"]
        # Dismissing a suggestion that no longer fits marks it out of date; one that fits again is rejected.
        new_status = "stale" if dismiss and body.count(s["original"]) != 1 else "rejected"
        changed = conn.execute("UPDATE write_suggestions SET status = ? WHERE id = ? AND status = 'pending'",
                               (new_status, s["id"])).rowcount
        status = conn.execute("SELECT status FROM write_suggestions WHERE id = ?", (s["id"],)).fetchone()["status"]
        place = _next_place(conn, batch_id, s["number"])
    if changed and new_status == "stale":
        _say(request, here, f"suggestion {s['number']} dismissed",
             f"Suggestion {s['number']} dismissed. Your text is unchanged.", "success", place)
    elif changed:
        _say(request, here, f"suggestion {s['number']} rejected",
             f"Suggestion {s['number']} rejected. Your text is unchanged.", "success", place)
    else:
        _say(request, here, f"suggestion {s['number']} was already decided",
             f"Suggestion {s['number']} was already decided: {STATUS_TEXT[status]}", "info", place)
    return redirect(_back_to(doc_id, batch_id, place))


def _join_numbers(numbers: list[int]) -> str:
    words = [str(n) for n in numbers]
    if len(words) <= 1:
        return "".join(words)
    return ", ".join(words[:-1]) + " and " + words[-1]


@router.post("/{doc_id}/suggestions/{batch_id}/accept-all")
async def accept_all(request: Request, doc_id: int, batch_id: int, user: dict = Depends(require_user)):
    await checked_form(request)
    doc = _owned_doc(doc_id, user)
    _owned_batch(doc, batch_id)
    accepted: list[int] = []
    stale: list[int] = []
    version = 0
    with db.session() as conn:
        _begin_write(conn)
        body = conn.execute("SELECT body FROM write_docs WHERE id = ?", (doc_id,)).fetchone()["body"]
        rows = conn.execute("SELECT * FROM write_suggestions WHERE batch_id = ? AND status = 'pending' ORDER BY number",
                            (batch_id,)).fetchall()
        for s in rows:
            if body.count(s["original"]) == 1:  # an earlier change may have made a second copy
                body = body.replace(s["original"], s["proposed"], 1)
                conn.execute("UPDATE write_suggestions SET status = 'accepted' WHERE id = ?", (s["id"],))
                accepted.append(s["number"])
            else:
                conn.execute("UPDATE write_suggestions SET status = 'stale' WHERE id = ?", (s["id"],))
                stale.append(s["number"])
        if accepted:
            note = (f"Accepted suggestion {accepted[0]}" if len(accepted) == 1
                    else f"Accepted suggestions {_join_numbers(accepted)}")
            _set_body(conn, doc_id, body, note)
            version = conn.execute("SELECT COUNT(*) AS n FROM write_versions WHERE doc_id = ?", (doc_id,)).fetchone()["n"]
    parts = []
    if accepted:
        parts.append(f"Accepted {plural(len(accepted), 'suggestion')} and saved your text as version {version}.")
    if stale:
        parts.append(f"{plural(len(stale), 'suggestion')} no longer matched your text, so "
                     f"{'it was' if len(stale) == 1 else 'they were'} not used: number {_join_numbers(stale)}.")
    if not parts:
        parts.append("There were no suggestions waiting for a decision. Nothing was changed.")
    outcomes = []
    if accepted:
        outcomes.append(f"{plural(len(accepted), 'suggestion')} accepted")
    if stale:
        outcomes.append(f"{len(stale)} out of date")
    here = f"/write/{doc_id}/suggestions/{batch_id}"
    _say(request, here, ", ".join(outcomes) if outcomes else "nothing was waiting, nothing changed",
         " ".join(parts), "success" if accepted else "info")
    return redirect(_back_to(doc_id, batch_id, 0))


@router.post("/{doc_id}/suggestions/{batch_id}/reject-all")
async def reject_all(request: Request, doc_id: int, batch_id: int, user: dict = Depends(require_user)):
    await checked_form(request)
    doc = _owned_doc(doc_id, user)
    _owned_batch(doc, batch_id)
    with db.session() as conn:
        cur = conn.execute("UPDATE write_suggestions SET status = 'rejected' WHERE batch_id = ? AND status = 'pending'",
                           (batch_id,))
        count = cur.rowcount
    here = f"/write/{doc_id}/suggestions/{batch_id}"
    if count:
        _say(request, here, f"{plural(count, 'suggestion')} rejected",
             f"Rejected {plural(count, 'suggestion')}. Your text is unchanged.")
    else:
        _say(request, here, "nothing was waiting, nothing changed",
             "There were no suggestions waiting for a decision. Nothing was changed.", "info")
    return redirect(_back_to(doc_id, batch_id, 0))


# ----- versions --------------------------------------------------------------------------------

def _clock(iso: str | None) -> str:
    return f"{iso[11:16]} UTC" if iso and len(iso) >= 16 and iso[13:14] == ":" else ""


@router.get("/{doc_id}/versions")
def versions_page(request: Request, doc_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    versions = list(reversed(_versions(doc_id)))
    for i, v in enumerate(versions):
        v["current"] = i == 0 and v["body"] == doc["body"]
        name = f"Version {v['number']}: {_lower_first(v['note'])}" if v["note"] else f"Version {v['number']}"
        clock = _clock(v["created_at"])
        v["link"] = name + (f", {clock}" if clock else "") + (" (current text)" if v["current"] else "")
    return render(request, "write/versions.html", title=f"Versions of {doc['title']}", doc=_doc_view(doc),
                  versions=versions, days=_versions_by_day(versions))


def _owned_version(doc: dict, version_id: int) -> dict:
    """One version of this document, with its number, without loading every other version's text."""
    with db.session() as conn:
        row = conn.execute("SELECT * FROM write_versions WHERE id = ? AND doc_id = ?", (version_id, doc["id"])).fetchone()
        if row is None:
            raise _not_found("version")
        number = conn.execute("SELECT COUNT(*) AS n FROM write_versions WHERE doc_id = ? AND id <= ?",
                              (doc["id"], version_id)).fetchone()["n"]
    v = dict(row)
    v.update(number=number, when=when(v["created_at"]), words=word_count(v["body"]))
    v["word_text"] = plural(v["words"], "word")
    v["saved"] = saved_sentence(v["note"], v["created_at"], prefix="Saved")
    return v


MAX_DIFFERENCES = 30
_DIFF_TEXT_MAX = 600


def _paragraphs(text: str) -> list[str]:
    """Paragraphs (blocks between blank lines), each on one line, without Markdown heading marks."""
    out = []
    for block in re.split(r"\n[ \t]*\n", text or ""):
        words = " ".join(_MD_HEADING.sub("", ln) for ln in block.split("\n")).split()
        if words:
            out.append(" ".join(words))
    return out


def _short(text: str) -> str:
    return text if len(text) <= _DIFF_TEXT_MAX else text[:_DIFF_TEXT_MAX - 3].rstrip() + "..."


def version_differences(version_text: str, current_text: str) -> list[dict]:
    """The paragraphs that differ between a version and the current text, worked out by the site.
    Each is {"version": text or "", "current": text or ""}; a missing side means that text does not have it."""
    old, new = _paragraphs(version_text), _paragraphs(current_text)
    out = []
    for tag, i1, i2, j1, j2 in SequenceMatcher(None, old, new, autojunk=False).get_opcodes():
        if tag == "equal":
            continue
        out.append({"version": _short(" ".join(old[i1:i2])), "current": _short(" ".join(new[j1:j2]))})
    return out


@router.get("/{doc_id}/versions/{version_id}")
def version_page(request: Request, doc_id: int, version_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    v = _owned_version(doc, version_id)
    same = v["body"] == doc["body"]
    differences = [] if same else version_differences(v["body"], doc["body"])
    return render(request, "write/version.html", title=f"Version {v['number']} of {doc['title']}", doc=_doc_view(doc),
                  version=v, same_as_current=same, differences=differences[:MAX_DIFFERENCES],
                  more_differences=max(0, len(differences) - MAX_DIFFERENCES))


@router.get("/{doc_id}/versions/{version_id}/restore")
def back_to_version(doc_id: int, version_id: int, user: dict = Depends(require_user)):
    _owned_version(_owned_doc(doc_id, user), version_id)
    return redirect(f"/write/{doc_id}/versions/{version_id}")


@router.post("/{doc_id}/versions/{version_id}/restore")
async def restore(request: Request, doc_id: int, version_id: int, user: dict = Depends(require_user)):
    await checked_form(request)
    doc = _owned_doc(doc_id, user)
    v = _owned_version(doc, version_id)
    here = f"/write/{doc_id}"
    with db.session() as conn:
        # Under the write lock, so pressing Restore twice does not save the same versions twice.
        _begin_write(conn)
        current = conn.execute("SELECT body FROM write_docs WHERE id = ?", (doc_id,)).fetchone()["body"]
        same = v["body"] == current
        if not same:
            last = conn.execute("SELECT body FROM write_versions WHERE doc_id = ? ORDER BY id DESC LIMIT 1",
                                (doc_id,)).fetchone()
            if last is None or last["body"] != current:
                _add_version(conn, doc_id, current, f"Your text before restoring version {v['number']}")
            _set_body(conn, doc_id, v["body"], f"Restored version {v['number']}")
    if same:
        flash(request, f"Version {v['number']} is the same as your current text. Nothing was changed.", "info")
        _say(request, here, f"{doc['title']}: version {v['number']} is already your text")
    else:
        flash(request, f"Restored version {v['number']}. Your text from before is kept in the versions list.", "success")
        _say(request, here, f"{doc['title']}: version {v['number']} restored")
    return redirect(here)


# ----- export ----------------------------------------------------------------------------------

EXPORT_FORMATS = ("docx", "md", "txt", "html", "tex")


@router.get("/{doc_id}/export")
def export_page(request: Request, doc_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    slug = _slug(doc["title"])
    return render(request, "write/export.html", title=f"Export {doc['title']}", doc=_doc_view(doc),
                  files={ext: f"{slug}.{ext}" for ext in EXPORT_FORMATS})


def _body_without_title(doc: dict) -> str:
    """The text, minus a first line that only repeats the title as a heading (the export adds the title itself)."""
    lines = doc["body"].split("\n")
    for i, line in enumerate(lines):
        if not line.strip():
            continue
        m = re.match(r"^#\s+(.*?)\s*#*\s*$", line)
        if m and m.group(1).strip().lower() == doc["title"].strip().lower():
            return "\n".join(lines[i + 1:]).lstrip("\n")
        break
    return doc["body"]


def _heading(line: str) -> tuple[int, str] | None:
    m = re.match(r"^(#{1,6})\s+(.*?)\s*#*\s*$", line)
    if m and m.group(2):
        return len(m.group(1)), m.group(2)
    return None


def _export_lines(doc: dict) -> list[tuple[int, str]]:
    """The text under the title, one (level, text) per line: level 0 for ordinary lines, otherwise the heading's
    rank among the levels the member used (1 is the highest). So '## Section' then '#### Part' export as
    section and subsection, whatever marks were used, and no heading level is skipped."""
    lines = _body_without_title(doc).split("\n")
    headings = [_heading(line.strip()) for line in lines]
    rank = {level: i + 1 for i, level in enumerate(sorted({h[0] for h in headings if h}))}
    return [(rank[h[0]], h[1]) if h else (0, line) for line, h in zip(lines, headings)]


def export_markdown(doc: dict) -> str:
    return f"# {doc['title']}\n\n{_body_without_title(doc).rstrip()}\n"


def export_text(doc: dict) -> str:
    """Plain text: the title, then the text with each heading on a line of its own and no # marks."""
    out = [doc["title"], ""]
    for level, text in _export_lines(doc):
        if level:
            if out[-1] != "":
                out.append("")
            out += [text, ""]
        elif text.strip() or out[-1] != "":
            out.append(text.rstrip())
    return "\n".join(out).rstrip() + "\n"


def export_html(doc: dict, lang: str = "en") -> str:
    from markupsafe import escape
    lang = lang if _LANG_RE.match(lang or "") else "en"
    content = nest_headings(markdown(_body_without_title(doc)), 2)
    return (
        "<!doctype html>\n"
        f'<html lang="{escape(lang)}">\n<head>\n<meta charset="utf-8">\n'
        '<meta name="viewport" content="width=device-width, initial-scale=1">\n'
        f"<title>{escape(doc['title'])}</title>\n"
        "<style>body{max-width:45rem;margin:0 auto;padding:1rem;font-family:system-ui,sans-serif;"
        "font-size:1.125rem;line-height:1.6}</style>\n"
        "</head>\n<body>\n<main>\n"
        f"<h1>{escape(doc['title'])}</h1>\n{content}</main>\n</body>\n</html>\n"
    )


def export_docx(doc: dict) -> bytes:
    from docx import Document

    # Text pasted from other programs can hold control characters a Word file cannot contain.
    doc = dict(doc, title=" ".join(_xml_safe(doc["title"]).split()), body=_xml_safe(doc["body"]))
    document = Document()
    document.core_properties.title = doc["title"]
    document.add_heading(doc["title"], level=1)
    for level, line in _export_lines(doc):
        stripped = line.strip()
        if not stripped:
            continue
        if level:
            document.add_heading(stripped, level=min(level + 1, 9))
        elif re.match(r"^[-*]\s+", stripped):
            document.add_paragraph(re.sub(r"^[-*]\s+", "", stripped), style="List Bullet")
        else:
            document.add_paragraph(line.rstrip())
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


_LATEX_SPECIAL = {
    "\\": r"\textbackslash{}",
    "&": r"\&",
    "%": r"\%",
    "$": r"\$",
    "#": r"\#",
    "_": r"\_",
    "{": r"\{",
    "}": r"\}",
    "~": r"\textasciitilde{}",
    "^": r"\textasciicircum{}",
}
_LATEX_RE = re.compile("|".join(re.escape(k) for k in _LATEX_SPECIAL))


def latex_escape(text: str) -> str:
    return _LATEX_RE.sub(lambda m: _LATEX_SPECIAL[m.group(0)], text)


def export_tex(doc: dict) -> str:
    sections = {1: "section", 2: "subsection", 3: "subsubsection", 4: "paragraph", 5: "subparagraph", 6: "subparagraph"}
    out = [
        r"\documentclass{article}",
        r"\usepackage[utf8]{inputenc}",
        r"\usepackage[T1]{fontenc}",
        r"\title{" + latex_escape(doc["title"]) + "}",
        r"\date{}",
        r"\begin{document}",
        r"\maketitle",
        "",
    ]
    for level, line in _export_lines(doc):
        stripped = line.strip()
        if not stripped:
            continue
        if level:
            out.append("\\" + sections[level] + "{" + latex_escape(stripped) + "}")
        else:
            out.append(latex_escape(line.rstrip()))
        out.append("")
    out.append(r"\end{document}")
    return "\n".join(out) + "\n"


def _download(content: str | bytes, media_type: str, filename: str) -> Response:
    return Response(content=content, media_type=media_type,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@router.get("/{doc_id}/export.md")
def export_md_file(doc_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    return _download(export_markdown(doc), "text/markdown; charset=utf-8", _slug(doc["title"]) + ".md")


@router.get("/{doc_id}/export.txt")
def export_txt_file(doc_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    return _download(export_text(doc), "text/plain; charset=utf-8", _slug(doc["title"]) + ".txt")


@router.get("/{doc_id}/export.html")
def export_html_file(doc_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    return _download(export_html(doc, user.get("language") or "en"), "text/html; charset=utf-8",
                     _slug(doc["title"]) + ".html")


@router.get("/{doc_id}/export.docx")
def export_docx_file(doc_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    return _download(export_docx(doc), "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                     _slug(doc["title"]) + ".docx")


@router.get("/{doc_id}/export.tex")
def export_tex_file(doc_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    return _download(export_tex(doc), "application/x-tex; charset=utf-8", _slug(doc["title"]) + ".tex")


# ----- where to publish ------------------------------------------------------------------------

@router.post("/{doc_id}/venues")
async def venues(request: Request, doc_id: int, user: dict = Depends(require_user)):
    doc, kept, errors, saved = await _save_first(request, doc_id, user)
    if not errors:
        if not doc["body"].strip():
            errors["body"] = ("Your document has no text yet. Write or paste it in \"Your text\", so the AI knows what "
                              "it is about. Then choose \"Save and search the web for places to publish\".")
        errors.update(_stored_secret_errors(doc))
    if errors:
        return _ai_problems(request, doc, kept, errors, saved)
    return _start_job(request, user, doc, "write.venues", f"Where to publish: {doc['title']}", _venues_job, doc_id,
                      doc["created_at"], saved=saved)


def _venues_job(ctx, ai, doc_id: int, created_at: str) -> str:
    doc = _job_doc(ctx, doc_id, created_at)
    if doc is None:
        raise AIError(GONE)
    ctx.progress("The AI is searching the web for places to publish")
    excerpt = doc["body"][:VENUE_EXCERPT]
    prompt = (
        f"Find suitable places to publish this {KIND_NAMES.get(doc['kind'], doc['kind']).lower()}: journals, magazines, "
        "websites or preprint servers.\n"
        + (f"The writer is considering: {doc['target']}. Include it and say how well it fits.\n" if doc["target"] else "")
        + f"Title: {doc['title']}\nLength: about {word_count(doc['body'])} words.\n"
        f"The text (it may be cut short):\n<text>\n{excerpt}\n</text>\n\n"
        "Start with a short summary naming the best two or three choices and why. Then give each place its own "
        "level-1 Markdown heading (# Name), with:\n"
        "- what it publishes and who reads it;\n"
        "- submission requirements: length, format, referencing style, anonymised copy, and a link to the author guidelines;\n"
        "- fees: any submission fee or open access fee (article processing charge), with the amount and currency, "
        "or say clearly that you did not find one;\n"
        "- the submission system it uses (for example a named online system or email), and anything published about how "
        "well it works with screen readers. If you found nothing, write 'Nothing found about screen reader access'. Never guess.\n"
        "Only state facts you found in your sources. Say when a fact may be out of date."
    )
    result = ai.research(system=VENUES_SYSTEM, prompt=prompt, purpose="write.venues")
    text = getattr(result, "text", None)
    if not isinstance(text, str) or not text.strip():
        raise AIError("The AI searched but did not write an answer, so nothing was saved. Please try again.")
    sources = _clean_sources(getattr(result, "sources", None))
    with db.session() as conn:
        _begin_write(conn)
        if _get_doc(conn, ctx, doc_id, created_at) is None:
            raise AIError(GONE_DURING)
        cur = conn.execute(
            "INSERT INTO write_research (doc_id, kind, content_md, sources_json, created_at) VALUES (?, 'venues', ?, ?, ?)",
            (doc_id, text.strip(), json.dumps(sources), db.now()))
        rid = cur.lastrowid
    return f"/write/{doc_id}/publish/{rid}"


@router.get("/{doc_id}/publish/{research_id}")
def research_page(request: Request, doc_id: int, research_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    with db.session() as conn:
        row = conn.execute("SELECT * FROM write_research WHERE id = ? AND doc_id = ?", (research_id, doc_id)).fetchone()
    if row is None:
        raise _not_found("search result")
    research = dict(row)
    try:
        sources = _clean_sources(json.loads(research["sources_json"] or "[]"))
    except (ValueError, TypeError):
        sources = []
    return render(request, "write/research.html", title=f"Where to publish: {doc['title']}", doc=_doc_view(doc),
                  research=dict(research, when=when(research["created_at"]), day=date_words(research["created_at"])),
                  sources=sources)


# ----- cover letter ----------------------------------------------------------------------------

@router.post("/{doc_id}/cover-letter")
async def cover_letter(request: Request, doc_id: int, user: dict = Depends(require_user)):
    doc, kept, errors, saved = await _save_first(request, doc_id, user)
    notes = kept.get("notes", "")
    # Each message is attached to the field that fixes it, so the error summary link goes to a real field.
    if not errors:
        if not doc["body"].strip():
            errors["body"] = ("Your document has no text yet. Write or paste it in \"Your text\", so the letter can "
                              "describe it. Then choose \"Save and draft a cover letter\".")
        if not doc["target"]:
            errors["target"] = ("Fill in the publication you are sending this to, for example a journal or magazine. "
                                "Then choose \"Save and draft a cover letter\".")
        for field, problem in _stored_secret_errors(doc).items():
            errors.setdefault(field, problem)
    if len(notes) > MAX_INSTRUCTIONS:
        errors["notes"] = f"This is {len(notes):,} characters long. Shorten it to {MAX_INSTRUCTIONS:,} characters or fewer."
    elif _secret_error(notes):
        errors["notes"] = _secret_error(notes)
    if errors:
        return _ai_problems(request, doc, kept, errors, saved)
    return _start_job(request, user, doc, "write.cover_letter", f"Cover letter for {doc['title']}", _cover_letter_job,
                      doc_id, doc["created_at"], notes, saved=saved)


def _cover_letter_job(ctx, ai, doc_id: int, created_at: str, notes: str) -> str:
    doc = _job_doc(ctx, doc_id, created_at)
    if doc is None:
        raise AIError(GONE)
    ctx.progress("The AI is drafting your cover letter")
    prompt = (
        f"Draft a cover letter to the editor of {doc['target']}, submitting the "
        f"{KIND_NAMES.get(doc['kind'], doc['kind']).lower()} titled \"{doc['title']}\" "
        f"(about {word_count(doc['body'])} words).\n"
        + (f"The writer wants the letter to mention: {notes}\n" if notes else "")
        + "Use only facts found in the text or in the writer's notes. Never invent the writer's name, credentials, "
        "past publications, funding or affiliations. Never mention that the writer is blind or has low vision, or how "
        "they read, unless their text or notes ask you to. Put anything you do not know in square brackets, for example "
        "[Your name] or [Editor's name], so the writer can find and fill it in.\n"
        "Write plain text only: no Markdown, no headings, no emoji. Keep it under 400 words. "
        "Say briefly what the work is, why it suits this publication, and confirm it is not under consideration "
        "elsewhere as a bracketed item the writer must confirm.\n"
        f"The text:\n<text>\n{doc['body'][:MAX_AI_BODY]}\n</text>"
    )
    draft = ai.text(system=COVER_LETTER_SYSTEM, prompt=prompt, purpose="write.cover_letter")
    if not isinstance(draft, str) or not draft.strip():
        raise AIError("The AI did not write a letter, so no new document was made. Please try again.")
    draft = draft.strip().replace("\r\n", "\n")[:MAX_BODY]
    title = f"{COVER_PREFIX}{doc['title']}"[:MAX_TITLE]
    now = db.now()
    with db.session() as conn:
        _begin_write(conn)
        if _get_doc(conn, ctx, doc_id, created_at) is None:
            raise AIError(GONE_DURING)
        cur = conn.execute(
            "INSERT INTO write_docs (user_id, title, kind, body, target, created_at, updated_at) VALUES (?, ?, 'letter', ?, ?, ?, ?)",
            (doc["user_id"], title, draft, doc["target"], now, now))
        new_id = cur.lastrowid
        _add_version(conn, new_id, draft, DRAFTED_NOTE)
    return f"/write/{new_id}"


# ----- delete ----------------------------------------------------------------------------------

@router.get("/{doc_id}/delete")
def delete_page(request: Request, doc_id: int, user: dict = Depends(require_user)):
    doc = _owned_doc(doc_id, user)
    return render(request, "write/delete.html", title=f"Delete {doc['title']}?", doc=_doc_view(doc),
                  version_count=_version_summary(doc_id)[0])


@router.post("/{doc_id}/delete")
async def delete(request: Request, doc_id: int, user: dict = Depends(require_user)):
    await checked_form(request)
    doc = _owned_doc(doc_id, user)
    with db.session() as conn:
        conn.execute("DELETE FROM write_suggestions WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM write_batches WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM write_versions WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM write_research WHERE doc_id = ?", (doc_id,))
        conn.execute("DELETE FROM write_docs WHERE id = ? AND user_id = ?", (doc_id, user["id"]))
        jobs.forget_item(conn, user["id"], _item_ref(doc))
    flash(request, f"Deleted \"{doc['title']}\" and all its versions.", "success")
    _say(request, "/write", f"{doc['title']} deleted")
    return redirect("/write")


def forget_user(user_id: int) -> None:
    """Called when a member deletes their account: remove their documents (versions, suggestions cascade)."""
    with db.session() as conn:
        conn.execute("DELETE FROM write_docs WHERE user_id = ?", (user_id,))
